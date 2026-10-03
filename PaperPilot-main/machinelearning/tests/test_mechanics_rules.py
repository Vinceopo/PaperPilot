import io
import unittest

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from app.compliance import _page_token, run_compliance_scan
from app.documents import normalize_mechanics_rules, parse_document

BODY = (
    "This study examines how students format their manuscripts and how a checker can help them. "
    "It describes the problem, the method, and the expected results in plain language for readers."
)


def _page_field(paragraph):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def _set_number_format(section, fmt=None, start=None):
    element = section._sectPr.find(qn("w:pgNumType"))
    if element is None:
        element = OxmlElement("w:pgNumType")
        section._sectPr.append(element)
    if fmt:
        element.set(qn("w:fmt"), fmt)
    # add_section() copies the previous section's settings, including its restart.
    element.attrib.pop(qn("w:start"), None)
    if start is not None:
        element.set(qn("w:start"), str(start))


def _number_header(section, where="header", align=WD_ALIGN_PARAGRAPH.RIGHT, first_page=None):
    """first_page: None = same as other pages, 'hidden', or ('footer', WD_ALIGN_PARAGRAPH.CENTER)."""
    part = section.header if where == "header" else section.footer
    part.is_linked_to_previous = False
    paragraph = part.paragraphs[0]
    paragraph.alignment = align
    _page_field(paragraph)
    if first_page is not None:
        section.different_first_page_header_footer = True
        section.first_page_header.is_linked_to_previous = False
        section.first_page_footer.is_linked_to_previous = False
        if first_page != "hidden":
            first_where, first_align = first_page
            first_part = section.first_page_header if first_where == "header" else section.first_page_footer
            first_paragraph = first_part.paragraphs[0]
            first_paragraph.alignment = first_align
            _page_field(first_paragraph)


def _body(doc, text=BODY, count=2):
    for _ in range(count):
        paragraph = doc.add_paragraph(text)
        paragraph.paragraph_format.first_line_indent = Inches(0.5)
        paragraph.paragraph_format.line_spacing = 2.0


def _page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _heading(doc, text, bold=True):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(12)
    return paragraph


def _thesis(
    prelim_fmt="lowerRoman",
    body_start=1,
    title_first_page="hidden",
    chapter_first_page=("footer", WD_ALIGN_PARAGRAPH.CENTER),
    align=WD_ALIGN_PARAGRAPH.RIGHT,
):
    """Title page + 2 preliminary pages, then Chapter I (2 pages) and Chapter II, each chapter its own section."""
    doc = Document()
    _heading(doc, "A STUDY OF MANUSCRIPT FORMATTING")
    doc.add_paragraph("A thesis presented to the faculty")
    _page_break(doc)
    _heading(doc, "ABSTRACT")
    _body(doc)
    _page_break(doc)
    _heading(doc, "ACKNOWLEDGMENT")
    _body(doc)
    prelim = doc.sections[0]
    _number_header(prelim, align=align, first_page=title_first_page)
    if prelim_fmt:
        _set_number_format(prelim, fmt=prelim_fmt)

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    _heading(doc, "CHAPTER I")
    _heading(doc, "INTRODUCTION")
    _body(doc)
    _page_break(doc)
    _body(doc)
    _number_header(body, align=align, first_page=chapter_first_page)
    _set_number_format(body, fmt="decimal", start=body_start)

    second = doc.add_section(WD_SECTION.NEW_PAGE)
    _heading(doc, "CHAPTER II")
    _body(doc)
    _number_header(second, align=align, first_page=chapter_first_page)
    _set_number_format(second, fmt="decimal")
    buffer = io.BytesIO()
    doc.save(buffer)
    return parse_document(buffer.getvalue(), "docx")


PAGINATION_RULES = {
    "pagination": {
        "position": "Top right",
        "title_page": "Hidden but counted",
        "first_page_of_chapter": "Bottom center",
        "preliminary_style": "Lowercase Roman (i, ii, iii)",
        "body_numbering": "Arabic, restart at 1 on Chapter 1",
        "chapter_markers": ["chapter_roman"],
    }
}


def _issue_types(result, prefix=""):
    return {issue["issue_type"] for issue in result["issues"] if issue["issue_type"].startswith(prefix)}


def _section(result, name):
    return next(item for item in result["sections"] if item["section"] == name)


class PaginationTests(unittest.TestCase):
    def scan(self, parsed, rules):
        return run_compliance_scan(parsed, normalize_mechanics_rules(rules))

    def test_section_start_counts_as_new_page(self):
        parsed = _thesis()
        sections = [paragraph.get("section_index") for paragraph in parsed["paragraphs"] if paragraph["source"] == "body"]
        self.assertEqual(sections[0], 0)
        self.assertEqual(sections[-1], 2)
        self.assertEqual(parsed["sections"][1]["page_number"]["start"], 1)
        self.assertIsNone(parsed["sections"][2]["page_number"]["start"])
        self.assertEqual(parsed["sections"][0]["page_number"]["format"], "roman_lower")
        from app.compliance import _document_units

        pages = {unit["page_index"] for unit in _document_units(parsed, {}) if unit["text"].startswith("CHAPTER")}
        self.assertEqual(pages, {3, 5})

    def test_compliant_thesis_has_no_pagination_issues(self):
        result = self.scan(_thesis(), PAGINATION_RULES)
        self.assertEqual(_issue_types(result, "pagination"), set())
        self.assertEqual(_section(result, "Pagination")["formatting_score"], 100.0)
        self.assertGreater(_section(result, "Pagination")["issue_count"], -1)

    def test_title_page_number_shown_is_flagged(self):
        result = self.scan(_thesis(title_first_page=None), PAGINATION_RULES)
        self.assertIn("pagination_title_page", _issue_types(result))

    def test_left_position_is_checked(self):
        rules = {"pagination": {**PAGINATION_RULES["pagination"], "position": "Top left"}}
        result = self.scan(_thesis(), rules)
        self.assertIn("pagination_position", _issue_types(result))
        compliant = self.scan(_thesis(align=WD_ALIGN_PARAGRAPH.LEFT), rules)
        self.assertNotIn("pagination_position", _issue_types(compliant))

    def test_arabic_preliminary_pages_flagged_when_roman_required(self):
        result = self.scan(_thesis(prelim_fmt="decimal"), PAGINATION_RULES)
        self.assertIn("pagination_style", _issue_types(result))

    def test_body_must_restart_at_one(self):
        result = self.scan(_thesis(body_start=5), PAGINATION_RULES)
        issues = {issue["issue_type"]: issue for issue in result["issues"]}
        self.assertIn("pagination_sequence", issues)
        self.assertIn("restart at 1", issues["pagination_sequence"]["explanation"])

    def test_chapter_first_page_rules(self):
        hidden_rules = {"pagination": {**PAGINATION_RULES["pagination"], "first_page_of_chapter": "No page number shown"}}
        self.assertIn("pagination_chapter_first", _issue_types(self.scan(_thesis(), hidden_rules)))
        self.assertNotIn(
            "pagination_chapter_first", _issue_types(self.scan(_thesis(chapter_first_page="hidden"), hidden_rules))
        )
        same_rules = {"pagination": {**PAGINATION_RULES["pagination"], "first_page_of_chapter": "Same position as other pages"}}
        self.assertNotIn("pagination_position", _issue_types(self.scan(_thesis(chapter_first_page=None), same_rules)))

    def test_hidden_and_not_counted_title_page(self):
        rules = {"pagination": {**PAGINATION_RULES["pagination"], "title_page": "Hidden and not counted"}}
        self.assertIn("pagination_sequence", _issue_types(self.scan(_thesis(), rules)))

    def test_page_tokens(self):
        self.assertEqual(_page_token("iv"), (4, "roman_lower"))
        self.assertEqual(_page_token("Page 12 of 40"), (12, "arabic"))
        self.assertEqual(_page_token("- 7 -"), (7, "arabic"))
        self.assertIsNone(_page_token("iiii"))
        self.assertIsNone(_page_token("Introduction"))

    def test_pdf_page_numbers(self):
        def page(index, number, x=540, y=40, heading=None):
            lines = [{"text": heading or f"Body text line on page {index + 1} with enough words.",
                      "bbox": [72, 300, 540, 314], "spans": [{"text": "x", "size": 12, "font": "Times"}]}]
            if number is not None:
                lines.append({"text": number, "bbox": [x, y, x + 10, y + 12],
                              "spans": [{"text": number, "size": 12, "font": "Times"}]})
            for line_index, line in enumerate(lines):
                line.update(page_index=index, line_index=line_index)
            return {"page_index": index, "width_points": 612, "height_points": 792, "lines": lines}

        pages = [page(0, None), page(1, "ii"), page(2, "iii"),
                 page(3, None, heading="CHAPTER I"), page(4, "2"), page(5, "4")]
        parsed = {"text": "", "metadata": {"format": "pdf"}, "pages": pages}
        rules = {"pagination": {**PAGINATION_RULES["pagination"], "first_page_of_chapter": "No page number shown"}}
        result = self.scan(parsed, rules)
        issues = {issue["issue_type"]: issue for issue in result["issues"]}
        self.assertEqual(set(issues), {"pagination_sequence"})
        self.assertEqual(issues["pagination_sequence"]["locations"][0]["page"], 6)


class HeadingAndSpacingTests(unittest.TestCase):
    def _doc(self, heading_bold=True, body=BODY):
        doc = Document()
        heading = doc.add_paragraph("Introduction", style="Heading 1")
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.runs[0].bold = heading_bold
        _body(doc, body)
        buffer = io.BytesIO()
        doc.save(buffer)
        return parse_document(buffer.getvalue(), "docx")

    def test_heading_style(self):
        rules = normalize_mechanics_rules(
            {"font": {"heading_styles": {"heading1": {"bold": True, "case": "upper", "alignment": "center"}}}}
        )
        result = run_compliance_scan(self._doc(), rules)
        issue = next(issue for issue in result["issues"] if issue["issue_type"] == "heading1_style")
        self.assertIn("ALL CAPS", issue["explanation"])
        self.assertNotIn("bold", issue["explanation"])
        rules["font"]["heading_styles"]["heading1"]["case"] = "title"
        self.assertNotIn("heading1_style", _issue_types(run_compliance_scan(self._doc(), rules)))

    def test_word_spacing(self):
        rules = normalize_mechanics_rules({"word_spacing": "One space between words and after periods"})
        self.assertIn("word_spacing", _issue_types(run_compliance_scan(self._doc(body=BODY.replace(" how ", "  how ")), rules)))
        self.assertNotIn("word_spacing", _issue_types(run_compliance_scan(self._doc(), rules)))
        two = normalize_mechanics_rules({"word_spacing": "Two spaces after periods"})
        self.assertIn("word_spacing", _issue_types(run_compliance_scan(self._doc(), two)))
        self.assertNotIn(
            "word_spacing", _issue_types(run_compliance_scan(self._doc(body=BODY.replace(". It", ".  It")), two))
        )

    def test_landscape_pages(self):
        doc = Document()
        _body(doc)
        wide = doc.add_section(WD_SECTION.NEW_PAGE)
        wide.orientation = WD_ORIENT.LANDSCAPE
        wide.page_width, wide.page_height = wide.page_height, wide.page_width
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Year"
        buffer = io.BytesIO()
        doc.save(buffer)
        parsed = parse_document(buffer.getvalue(), "docx")
        base = {"paper": {"size": "8.5 x 11", "orientation": "Portrait"}}
        strict = run_compliance_scan(parsed, normalize_mechanics_rules(base))
        self.assertIn("paper_orientation", _issue_types(strict))
        allowed = {"paper": {**base["paper"], "landscape_pages": "Allowed for tables and figures"}}
        self.assertNotIn("paper_orientation", _issue_types(run_compliance_scan(parsed, normalize_mechanics_rules(allowed))))


class NormalizeTests(unittest.TestCase):
    def test_legacy_values_map_to_dropdown_choices(self):
        rules = normalize_mechanics_rules(
            {"pagination": {"position": "upper right corner", "first_page_of_chapter": "Counted but not numbered"}}
        )
        self.assertEqual(rules["pagination"]["position"], "Top right")
        self.assertEqual(rules["pagination"]["first_page_of_chapter"], "No page number shown")

    def test_new_fields_survive(self):
        rules = normalize_mechanics_rules(
            {
                **PAGINATION_RULES,
                "word_spacing": "Two spaces after periods",
                "paper": {"orientation": "Portrait", "landscape_pages": "Allowed on any page"},
                "font": {"heading_styles": {"heading2": {"bold": "Yes", "alignment": "Left"}}},
            }
        )
        self.assertEqual(rules["pagination"]["chapter_markers"], ["chapter_roman"])
        self.assertEqual(rules["pagination"]["title_page"], "Hidden but counted")
        self.assertEqual(rules["word_spacing"], "Two spaces after periods")
        self.assertEqual(rules["paper"]["landscape_pages"], "Allowed on any page")
        self.assertEqual(rules["font"]["heading_styles"], {"heading2": {"bold": True, "alignment": "left"}})


if __name__ == "__main__":
    unittest.main()
