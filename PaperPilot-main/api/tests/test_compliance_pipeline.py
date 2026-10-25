import copy
import unittest
from unittest.mock import patch

from app import compliance_db, emailer, otp, profiles
from app.compliance import run_compliance_scan
from app.documents import (
    DocumentError,
    _mechanics_rule_count,
    assert_mechanics_guide,
    build_sample_mechanics_docx,
    derive_mechanics_rules,
    normalize_mechanics_rules,
    parse_document,
    validate_document,
)


class FakeReference:
    def __init__(self, store, path="/"):
        self.store = store
        self.parts = [part for part in path.split("/") if part]

    def _parent(self, create=False):
        node = self.store
        for part in self.parts[:-1]:
            if create:
                node = node.setdefault(part, {})
            else:
                node = node.get(part, {})
        return node

    def get(self, shallow=False):
        node = self.store
        for part in self.parts:
            if not isinstance(node, dict) or part not in node:
                return None
            node = node[part]
        if shallow and isinstance(node, dict):
            return {key: True for key in node}
        return copy.deepcopy(node)

    def child(self, path):
        return FakeReference(self.store, "/" + "/".join([*self.parts, *path.split("/")]))

    def set(self, value):
        if not self.parts:
            self.store.clear()
            if isinstance(value, dict):
                self.store.update(copy.deepcopy(value))
            return
        parent = self._parent(create=True)
        parent[self.parts[-1]] = copy.deepcopy(value)

    def delete(self):
        if self.parts:
            self._parent().pop(self.parts[-1], None)

    def update(self, updates):
        for relative, value in updates.items():
            path = "/" + "/".join([*self.parts, *relative.split("/")])
            ref = FakeReference(self.store, path)
            ref.delete() if value is None else ref.set(value)

    def transaction(self, callback):
        current = self.get()
        updated = callback(current)
        if updated is None:
            self.delete()
        else:
            self.set(updated)
        return copy.deepcopy(updated)


class FirebaseTestCase(unittest.TestCase):
    def setUp(self):
        self.store = {}
        factory = lambda path: FakeReference(self.store, path)
        self.patchers = [
            patch.object(compliance_db, "_reference", side_effect=factory),
            patch.object(otp, "_reference", side_effect=factory),
            patch.object(profiles, "_reference", side_effect=factory),
        ]
        for patcher in self.patchers:
            patcher.start()
        otp._secret_cache = None
        self.old_secret = otp.settings.otp_secret
        otp.settings.otp_secret = "unit-test-secret"

    def tearDown(self):
        otp.settings.otp_secret = self.old_secret
        otp._secret_cache = None
        for patcher in reversed(self.patchers):
            patcher.stop()


class MechanicsExtractionTests(unittest.TestCase):
    def test_extracts_only_present_rules(self):
        rules = derive_mechanics_rules(
            "Use Times New Roman at 12 pt. Pages must use letter paper with 1 inch margins. "
            "Use double line spacing and APA citation style."
        )
        self.assertEqual(rules["font"]["families"], ["Times New Roman"])
        self.assertEqual(rules["font"]["sizes_points"], [12.0])
        self.assertEqual(rules["paper_size"]["name"], "LETTER")
        self.assertEqual(rules["line_spacing"], 2.0)
        self.assertEqual(rules["citation_style"], "APA")
        self.assertNotIn("first_line_indent_inches", rules)

    def test_prose_margins_and_typeface_size_from_outline_guide(self):
        rules = derive_mechanics_rules(
            "A. Typeface - Arial/Times New Roman font, size 12 must be used throughout the "
            "manuscript unless otherwise indicated (as in Titles, Headings)\n"
            "b. Spacing - Double spacing throughout the manuscript. Single spacing for "
            "bibliography entries and captions.\n"
            "C. Margins - 1.5 inch on the left, 1 inch on the right, top and bottom\n"
            "Preliminary Pages\n"
        )
        self.assertEqual(
            rules["margins_inches"],
            {"left": 1.5, "right": 1.0, "top": 1.0, "bottom": 1.0},
        )
        self.assertEqual(rules["font"]["families"][:2], ["Arial", "Times New Roman"])
        self.assertEqual(rules["font"]["body_size"], 12.0)
        self.assertEqual(rules["line_spacing"], 2.0)
        self.assertNotIn("size", rules.get("paper") or {})

    def test_one_side_and_other_sides_margins(self):
        rules = derive_mechanics_rules(
            "Margins: 1.5 inches on the left and 1 inch on all other sides."
        )
        self.assertEqual(
            rules["margins_inches"],
            {"left": 1.5, "right": 1.0, "top": 1.0, "bottom": 1.0},
        )

    def test_listed_sides_share_one_margin_value(self):
        rules = derive_mechanics_rules(
            "Page size should be 8.5 x 11-inches.\n"
            "All margins (left, right, top and bottom) should be 1.5 inches (3.8 cm), including your\n"
            "tables and figures.\nSingle space your text.\nFont:\n"
            "1. Main Body\u201412 pt. Times or the closest comparable font available\n"
            "2. Footnotes\u201410 pt. Times or the closest comparable font available\n"
        )
        self.assertEqual(
            rules["margins_inches"],
            {"left": 1.5, "right": 1.5, "top": 1.5, "bottom": 1.5},
        )
        self.assertEqual(rules["font"]["type"], "Times New Roman")
        self.assertEqual(rules["font"]["body_size"], 12.0)
        self.assertNotIn("table_layout_requirements", rules)
        self.assertNotIn("figure_layout_requirements", rules)

    def test_side_list_label_and_font_type_and_size(self):
        rules = derive_mechanics_rules(
            'b. Margins: Top, Right and Bottom: 1.00"\n'
            'Left: 1.25" to allow for binding\n'
            "c. Font Type and Size: Times New Roman; 12 points\n"
        )
        self.assertEqual(
            rules["margins_inches"],
            {"top": 1.0, "right": 1.0, "bottom": 1.0, "left": 1.25},
        )
        self.assertEqual(rules["font"]["body_size"], 12.0)

    def test_side_list_margins_with_curly_inch_marks(self):
        """PDF text often uses curly quotes as inch marks (Mapúa-style guides)."""
        rules = derive_mechanics_rules(
            "b. Margins: Top, Right and Bottom: 1.00\u201d\n"
            "Left: 1.25\u201d to allow for binding\n"
            "c. Font Type and Size: Times New Roman, 12 points\n"
        )
        self.assertEqual(
            rules["margins_inches"],
            {"top": 1.0, "right": 1.0, "bottom": 1.0, "left": 1.25},
        )

    def test_side_list_margins_when_inch_mark_missing(self):
        rules = derive_mechanics_rules(
            "b. Margins: Top, Right and Bottom: 1.00\n"
            "Left: 1.25 to allow for binding\n"
        )
        self.assertEqual(
            rules["margins_inches"],
            {"top": 1.0, "right": 1.0, "bottom": 1.0, "left": 1.25},
        )

    def test_outline_layout_for_tables_and_figures(self):
        rules = derive_mechanics_rules(
            "f. Pagination\n"
            "i. Top Right\n"
            "ii. No page shown on first page of every chapter\n"
            "g. Page Breaks: Do page break only when starting a new chapter.\n"
            "h. Layout for Tables\n"
            "Table number in bold, title in italic title case above the table\n"
            "i. Layout for Figures\n"
            "Figure number in bold with the title underlined below the image\n"
            "j. Citation Format: APA\n"
        )
        self.assertEqual(rules["pagination"]["position"], "Top right")
        self.assertEqual(rules["pagination"]["first_page_of_chapter"], "No page number shown")
        self.assertIn("chapter", rules["page_break_requirements"][0].lower())
        self.assertIn("table", rules["table_layout_requirements"][0].lower())
        self.assertIn("figure", rules["figure_layout_requirements"][0].lower())
        self.assertEqual(rules["citation_style"], "APA")

    def test_sample_mechanics_docx_fills_every_format_field(self):
        parsed = parse_document(build_sample_mechanics_docx(), "docx")
        rules = derive_mechanics_rules(parsed["text"])
        self.assertEqual(
            rules["paper"],
            {
                "size": "8.5 x 11",
                "orientation": "Portrait",
                "substance": "20",
                "landscape_pages": "Allowed for tables and figures",
            },
        )
        self.assertEqual(rules["line_spacing"], 2.0)
        self.assertEqual(rules["first_line_indent_inches"], 0.5)
        self.assertEqual(rules["alignment"], "left")
        self.assertEqual(
            rules["margins_inches"],
            {"top": 1.0, "bottom": 1.0, "left": 1.0, "right": 1.0, "gutter": 0.0, "header": 0.5, "footer": 0.5},
        )
        font = rules["font"]
        self.assertEqual(font["type"], "Times New Roman")
        self.assertEqual(font["color"], "Black/Automatic")
        self.assertEqual((font["heading1_size"], font["heading2_size"], font["body_size"]), (12.0, 12.0, 12.0))
        self.assertEqual(font["heading_levels"], 2)
        self.assertNotIn("heading3_size", font)
        self.assertEqual(rules["pagination"]["position"], "Top right")
        self.assertEqual(rules["pagination"]["first_page_of_chapter"], "Same position as other pages")
        self.assertEqual(
            rules["page_break_requirements"], ["Only to start each new chapter and the reference list"]
        )
        self.assertTrue(rules["table_layout_requirements"][0].endswith("both placed above it"))
        self.assertTrue(rules["figure_layout_requirements"][0].endswith("both placed above the image"))
        self.assertEqual(rules["citation_style"], "APA")

    def test_rejects_extension_signature_mismatch(self):
        with self.assertRaises(DocumentError):
            validate_document("paper.pdf", b"not a pdf", 100)

    def test_wrong_font_is_grouped_as_critical(self):
        parsed = {
            "text": "Example line",
            "metadata": {"format": "pdf", "pagination_fidelity": "fixed_source_pages"},
            "pages": [{
                "page_index": 1,
                "width_points": 612,
                "height_points": 792,
                "lines": [{
                    "line_index": 4,
                    "text": "Example line",
                    "spans": [{"text": "Example line", "font": "Arial", "size": 12, "bold": False}],
                }],
            }],
        }
        result = run_compliance_scan(
            parsed, {"font": {"families": ["Times New Roman"], "sizes_points": [12]}}, "free"
        )
        self.assertEqual(result["issues"][0]["severity"], "critical")
        loc0 = result["issues"][0]["locations"][0]
        self.assertEqual(loc0["page"], 2)
        self.assertEqual(loc0["line"], 1)
        self.assertEqual(loc0["section"], "Fonts")
        self.assertEqual(loc0["excerpt"], "Example line")
        self.assertEqual(loc0.get("page_width"), 612)
        self.assertEqual(loc0.get("page_height"), 792)
        self.assertIn("explanation", loc0)
        self.assertEqual([item["section"] for item in result["sections"]], [
            "Fonts", "Margins", "Indentation", "Spacing", "Alignment", "Pagination",
        ])
        fonts = next(item for item in result["sections"] if item["section"] == "Fonts")
        self.assertEqual(fonts["formatting_score"], 0)
        self.assertEqual(result["overall_score"], fonts["formatting_score"])

    def test_sample_guide_is_accepted_as_mechanics(self):
        text = parse_document(build_sample_mechanics_docx(), "docx")["text"]
        rules = derive_mechanics_rules(text)
        assert_mechanics_guide(text, rules)

    def test_manuscript_text_is_rejected_as_mechanics(self):
        manuscript = (
            "Chapter 1 Introduction\n"
            "Acknowledgements\n"
            "The completion of this research study would not have been possible without support. "
            * 80
            + "\nTable of Contents\nReview of Related Literature\nMethodology\n"
            "Findings and Discussion\nReferences\n"
        )
        rules = derive_mechanics_rules(manuscript)
        with self.assertRaises(DocumentError) as raised:
            assert_mechanics_guide(manuscript, rules)
        self.assertIn("manuscript", str(raised.exception).lower())

    def test_outline_school_guide_without_sample_titles_is_accepted(self):
        """Real school mechanics PDFs rarely use our sample section titles."""
        guide = (
            "Capstone Project Manuscript Mechanics\n"
            "Please refer to the following document formatting standards:\n"
            "a. Paper\n"
            "i. Size : 8.5 x 11\n"
            "ii. Orientation : Portrait\n"
            "iii. Substance : 20\n"
            "b. Spacing : Double\n"
            "c. Margins\n"
            "i. Top : 1 inch\n"
            "ii. Bottom : 1 inch\n"
            "iii. Left : 1.5 inches\n"
            "iv. Right : 1 inch\n"
            "d. Font\n"
            "i. Type : Times New Roman\n"
            "ii. Body text size : 12 pt\n"
            "iii. Heading 1 size : 14 pt\n"
            "e. Pagination : Top right\n"
            "f. Citation style : APA\n"
        )
        rules = derive_mechanics_rules(guide)
        assert_mechanics_guide(guide, rules)
        self.assertGreaterEqual(_mechanics_rule_count(rules), 3)

    def test_sample_guide_fills_every_mechanics_field(self):
        from app.documents import build_sample_mechanics_docx, normalize_mechanics_rules

        text = parse_document(build_sample_mechanics_docx(), "docx")["text"]
        rules = normalize_mechanics_rules(derive_mechanics_rules(text))
        self.assertEqual(rules["pagination"], {
            "position": "Top right",
            "first_page_of_chapter": "Same position as other pages",
            "title_page": "Number shown",
            "body_numbering": "Arabic, continuous",
            "chapter_markers": ["chapter_roman", "chapter_arabic"],
        })
        self.assertEqual(rules["font"]["heading_styles"]["heading1"], {"bold": True, "case": "title", "alignment": "center"})
        self.assertEqual(rules["word_spacing"], "One space between words and after periods")
        self.assertEqual(rules["paper"]["landscape_pages"], "Allowed for tables and figures")

    def test_overall_score_averages_checked_breakdown_metrics(self):
        parsed = {
            "text": "Line one\nLine two",
            "metadata": {"format": "pdf", "pagination_fidelity": "fixed_source_pages"},
            "pages": [{
                "page_index": 0,
                "width_points": 612,
                "height_points": 792,
                "lines": [
                    {
                        "line_index": 0,
                        "text": "Line one uses Arial here",
                        "spans": [{"text": "Line one uses Arial here", "font": "Arial", "size": 12, "bold": False}],
                        "bbox": [72, 72, 500, 86],
                    },
                    {
                        "line_index": 1,
                        "text": "Line two uses Times New Roman",
                        "spans": [{"text": "Line two uses Times New Roman", "font": "TimesNewRoman", "size": 12, "bold": False}],
                        "bbox": [72, 90, 500, 104],
                    },
                ],
            }],
        }
        result = run_compliance_scan(
            parsed,
            {
                "font": {"families": ["Times New Roman"], "sizes_points": [12]},
                "margins_inches": {"top": 1, "bottom": 1, "left": 1, "right": 1},
            },
            "free",
        )
        scores = {item["section"]: item["formatting_score"] for item in result["sections"]}
        self.assertEqual(scores["Fonts"], 50)
        self.assertEqual(scores["Margins"], 100)
        self.assertEqual(result["overall_score"], 75.0)
        self.assertEqual(result["issues"][0]["locations"][0]["page"], 1)
        self.assertEqual(result["issues"][0]["locations"][0]["line"], 1)
        self.assertEqual(result["page_count"], 1)

    def test_docx_line_numbers_restart_each_page(self):
        paragraphs = []
        for index in range(45):
            paragraphs.append({
                "paragraph_index": index,
                "style": "Normal",
                "text": f"Body paragraph number {index + 1} with enough words.",
                "lines": [{"line_index": index, "text": f"Body paragraph number {index + 1} with enough words."}],
                "runs": [{"text": f"Body paragraph number {index + 1} with enough words.", "font": "Arial", "size": 11}],
                "formatting": {"line_spacing": 1.5, "alignment": "left", "first_line_indent_inches": 0.0},
            })
        parsed = {
            "text": "\n".join(item["text"] for item in paragraphs),
            "metadata": {"format": "docx"},
            "paragraphs": paragraphs,
            "sections": [{
                "section_index": 0,
                "page_width_inches": 8.5,
                "page_height_inches": 11.0,
                "top_margin_inches": 1.0,
                "bottom_margin_inches": 1.0,
                "left_margin_inches": 1.0,
                "right_margin_inches": 1.0,
            }],
        }
        result = run_compliance_scan(
            parsed,
            {"font": {"families": ["Times New Roman"], "sizes_points": [11]}, "line_spacing": 1.5},
            "free",
        )
        locations = result["issues"][0]["locations"]
        self.assertGreater(result["page_count"], 1)
        self.assertEqual(locations[0]["page"], 1)
        self.assertEqual(locations[0]["line"], 1)
        page_two = [loc for loc in locations if loc["page"] == 2]
        self.assertTrue(page_two)
        self.assertEqual(min(loc["line"] for loc in page_two), 1)

    def test_heading_and_body_use_distinct_mechanics_sizes(self):
        parsed = {
            "text": "CHAPTER 1 INTRODUCTION\nThis body paragraph is written in eleven point text for the study.",
            "metadata": {"format": "pdf", "pagination_fidelity": "fixed_source_pages"},
            "pages": [{
                "page_index": 0,
                "width_points": 612,
                "height_points": 792,
                "lines": [
                    {
                        "line_index": 0,
                        "text": "CHAPTER 1 INTRODUCTION",
                        "spans": [{"text": "CHAPTER 1 INTRODUCTION", "font": "TimesNewRoman", "size": 14, "bold": True}],
                    },
                    {
                        "line_index": 1,
                        "text": "This body paragraph is written in eleven point text for the study.",
                        "spans": [{"text": "This body paragraph is written in eleven point text for the study.", "font": "TimesNewRoman", "size": 14, "bold": False}],
                    },
                ],
            }],
        }
        result = run_compliance_scan(
            parsed,
            {
                "font": {
                    "families": ["Times New Roman"],
                    "heading1_size": 14,
                    "heading2_size": 12,
                    "heading3_content_size": 11,
                    "sizes_points": [14, 12, 11],
                }
            },
            "free",
        )
        size_issues = [
            issue for issue in result["issues"] if str(issue["issue_type"]).startswith("font_size")
        ]
        self.assertTrue(size_issues)
        self.assertEqual(size_issues[0]["locations"][0]["line"], 2)

    def test_guide_heading_levels_follow_the_guide(self):
        rules = derive_mechanics_rules(
            "Font\nHeading 1 size: 16 pt\nHeading 2 size: 14 pt\nHeading 4 size: 11 pt\n"
            "Body text size: 12 pt\nHeading 4 style: Bold italic, Title Case, flush left\n"
        )
        font = rules["font"]
        self.assertEqual(font["heading_levels"], 4)
        self.assertEqual(
            (font["heading1_size"], font["heading2_size"], font["heading4_size"], font["body_size"]),
            (16.0, 14.0, 11.0, 12.0),
        )
        self.assertNotIn("heading3_size", font)
        self.assertEqual(font["heading_styles"]["heading4"], {"bold": True, "italic": True, "case": "title", "alignment": "left"})

        only_one = derive_mechanics_rules("Font\nHeading 1 size: 14 pt\nFont type: Arial\n")["font"]
        self.assertEqual(only_one["heading_levels"], 1)
        self.assertNotIn("heading_levels", derive_mechanics_rules("Font\nFont type: Arial\n").get("font", {}))

    def test_combined_heading3_and_content_size_sets_both(self):
        font = derive_mechanics_rules("Font\nHeading 1 size: 16 pt\nHeading 3 and content size: 12 pt\n")["font"]
        self.assertEqual((font["heading3_size"], font["body_size"]), (12.0, 12.0))
        self.assertEqual(font["heading_levels"], 3)

    def test_saved_rules_keep_every_heading_level(self):
        rules = normalize_mechanics_rules(
            {"font": {"heading5_size": 10, "body_size": 12, "heading_levels": 5,
                      "heading_styles": {"heading5": {"bold": "yes"}}}}
        )
        font = rules["font"]
        self.assertEqual((font["heading5_size"], font["body_size"], font["heading_levels"]), (10.0, 12.0, 5))
        self.assertEqual(font["heading3_content_size"], 12.0)
        self.assertEqual(font["heading_styles"], {"heading5": {"bold": True}})

    def test_heading_four_is_checked_against_its_own_rules(self):
        text = "Methods Overview"
        body = "This body paragraph is long enough to be ordinary text in the manuscript."
        parsed = {
            "text": f"{text}\n{body}",
            "metadata": {"format": "docx"},
            "paragraphs": [
                {
                    "paragraph_index": index,
                    "source": "body",
                    "style": style,
                    "text": value,
                    "lines": [{"line_index": 0, "text": value}],
                    "runs": [{"text": value, "font": "Times New Roman", "size": size, "bold": False}],
                    "formatting": {"line_spacing": 2.0, "alignment": "left"},
                }
                for index, (style, value, size) in enumerate((("Heading 4", text, 11), ("Normal", body, 12)))
            ],
            "sections": [{
                "section_index": 0,
                "page_width_inches": 8.5,
                "page_height_inches": 11.0,
                "top_margin_inches": 1.0,
                "bottom_margin_inches": 1.0,
                "left_margin_inches": 1.0,
                "right_margin_inches": 1.0,
            }],
        }
        rules = normalize_mechanics_rules({
            "font": {
                "families": ["Times New Roman"],
                "heading4_size": 12,
                "body_size": 12,
                "heading_levels": 4,
                "heading_styles": {"heading4": {"bold": True}},
            }
        })
        issues = run_compliance_scan(parsed, rules, "free")["issues"]
        titles = {issue["title"] for issue in issues}
        self.assertIn("Heading 4 font size differs", titles)
        self.assertIn("Heading 4 style differs", titles)
        self.assertNotIn("body text font size differs", titles)

    def _heading_doc(self, heading, *, style="Normal", size=12, bold=True, body=None):
        body = body or "This body paragraph is long enough to be ordinary text in the manuscript."
        return {
            "text": f"{heading}\n{body}",
            "metadata": {"format": "docx"},
            "paragraphs": [
                {
                    "paragraph_index": index,
                    "source": "body",
                    "style": para_style,
                    "text": value,
                    "lines": [{"line_index": 0, "text": value}],
                    "runs": [{"text": value, "font": "Times New Roman", "size": para_size, "bold": para_bold}],
                    "formatting": {"line_spacing": 2.0, "alignment": "left"},
                }
                for index, (para_style, value, para_size, para_bold) in enumerate(
                    ((style, heading, size, bold), ("Normal", body, 12, False))
                )
            ],
            "sections": [{
                "section_index": 0,
                "page_width_inches": 8.5,
                "page_height_inches": 11.0,
                "top_margin_inches": 1.0,
                "bottom_margin_inches": 1.0,
                "left_margin_inches": 1.0,
                "right_margin_inches": 1.0,
            }],
        }

    def test_heading_one_only_does_not_invent_heading_two(self):
        parsed = self._heading_doc("Related Work", style="Heading 1", bold=True)
        rules = normalize_mechanics_rules({
            "font": {
                "families": ["Times New Roman"],
                "heading1_size": 12,
                "body_size": 12,
                "heading_levels": 1,
                "heading_styles": {"heading1": {"bold": True, "case": "upper"}},
            }
        })
        types = {issue["issue_type"] for issue in run_compliance_scan(parsed, rules, "free")["issues"]}
        self.assertIn("heading1_style", types)
        self.assertNotIn("heading2_style", types)

    def test_bold_body_at_shared_size_is_not_heading1(self):
        """Bold Normal paragraphs must stay body when heading and body sizes match."""
        parsed = self._heading_doc(
            "The completion of this research study would not have been possible.",
            style="Normal",
            size=12,
            bold=True,
            body="Ordinary following paragraph with enough words for the scan.",
        )
        # Shorten first line to the heuristic length window while remaining prose.
        parsed["paragraphs"][0]["text"] = "Support from advisers made this possible."
        parsed["paragraphs"][0]["lines"][0]["text"] = "Support from advisers made this possible."
        parsed["paragraphs"][0]["runs"][0]["text"] = "Support from advisers made this possible."
        rules = normalize_mechanics_rules({
            "font": {
                "families": ["Times New Roman"],
                "heading1_size": 12,
                "body_size": 12,
                "heading_levels": 1,
            }
        })
        titles = {issue["title"] for issue in run_compliance_scan(parsed, rules, "free")["issues"]}
        self.assertFalse(any("Heading 1" in title for title in titles))

    def test_no_heading_rules_do_not_flag_heading_issues(self):
        parsed = self._heading_doc("Related Work", style="Heading 2", bold=True)
        rules = normalize_mechanics_rules({"font": {"families": ["Times New Roman"], "body_size": 12}})
        issues = run_compliance_scan(parsed, rules, "free")["issues"]
        types = {issue["issue_type"] for issue in issues}
        titles = {issue["title"] for issue in issues}
        self.assertFalse(any(item.endswith("_style") and item.startswith("heading") for item in types))
        self.assertFalse(any("Heading" in title and "font size" in title for title in titles))

    def test_word_heading_outside_the_format_is_not_that_level(self):
        parsed = self._heading_doc("Methods Overview", style="Heading 5", size=10, bold=True)
        rules = normalize_mechanics_rules({
            "font": {
                "families": ["Times New Roman"],
                "heading1_size": 14,
                "heading2_size": 12,
                "body_size": 12,
                "heading_levels": 2,
            }
        })
        titles = {issue["title"] for issue in run_compliance_scan(parsed, rules, "free")["issues"]}
        self.assertNotIn("Heading 5 font size differs", titles)

    def test_segoe_ui_emoji_does_not_override_times_new_roman(self):
        parsed = {
            "text": "This paragraph is Times New Roman with an emoji glyph nearby.",
            "metadata": {"format": "docx"},
            "paragraphs": [{
                "paragraph_index": 0,
                "source": "body",
                "style": "Normal",
                "text": "This paragraph is Times New Roman with an emoji glyph nearby.",
                "lines": [{
                    "line_index": 0,
                    "text": "This paragraph is Times New Roman with an emoji glyph nearby.",
                }],
                "runs": [
                    {"text": "✅", "font": "Segoe UI Emoji", "size": 11},
                    {
                        "text": "This paragraph is Times New Roman with an emoji glyph nearby.",
                        "font": "Times New Roman",
                        "size": 11,
                    },
                ],
                "formatting": {"line_spacing": 1.5, "alignment": "left"},
            }],
            "sections": [{
                "section_index": 0,
                "page_width_inches": 8.5,
                "page_height_inches": 11.0,
                "top_margin_inches": 1.0,
                "bottom_margin_inches": 1.0,
                "left_margin_inches": 1.0,
                "right_margin_inches": 1.0,
            }],
        }
        result = run_compliance_scan(
            parsed,
            {"font": {"families": ["Times New Roman"], "sizes_points": [11]}},
            "free",
        )
        family_issues = [issue for issue in result["issues"] if issue["issue_type"] == "font_family"]
        self.assertEqual(family_issues, [])
        scores = {item["section"]: item["formatting_score"] for item in result["sections"]}
        self.assertEqual(scores["Fonts"], 100.0)

    def test_visual_typeable_lines_wrap_and_restart_each_page(self):
        from app.compliance import _document_units, _soft_wrap_visual_lines

        wrapped = _soft_wrap_visual_lines(
            "This is a long academic sentence that should wrap across multiple typeable lines on the manuscript page.",
            40,
        )
        self.assertGreater(len(wrapped), 1)
        self.assertTrue(all(len(line) <= 40 for line in wrapped))

        long_text = (
            "This chapter explains the complete methodology of the study including participants instruments "
            "procedures and the statistical treatment of data for the BSIT manuscript format compliance tool."
        )
        paragraphs = []
        for index in range(8):
            paragraphs.append({
                "paragraph_index": index,
                "source": "body",
                "style": "Normal",
                "text": long_text,
                "lines": [{"line_index": index, "text": long_text}],
                "runs": [{"text": long_text, "font": "Times New Roman", "size": 11}],
                "formatting": {"line_spacing": 1.5, "alignment": "left"},
            })
        parsed = {
            "text": "\n".join(p["text"] for p in paragraphs),
            "metadata": {"format": "docx"},
            "paragraphs": paragraphs,
            "sections": [{
                "section_index": 0,
                "page_width_inches": 8.5,
                "page_height_inches": 11.0,
                "top_margin_inches": 1.0,
                "bottom_margin_inches": 1.0,
                "left_margin_inches": 1.0,
                "right_margin_inches": 1.0,
            }],
        }
        rules = {
            "font": {"families": ["Times New Roman"], "heading3_content_size": 11, "sizes_points": [11]},
            "line_spacing": 1.5,
            "margins_inches": {"top": 1, "bottom": 1, "left": 1, "right": 1},
        }
        units = _document_units(parsed, rules)
        self.assertGreater(len(units), 8)  # wrapped into more visual lines than paragraphs
        by_page = {}
        for unit in units:
            if (unit.get("source") or "body") in {"header", "footer"}:
                continue
            by_page.setdefault(unit["page_index"], []).append(unit["line_index"])
        self.assertIn(0, by_page)
        self.assertEqual(min(by_page[0]), 0)  # page line 1 in UI
        if 1 in by_page:
            self.assertEqual(min(by_page[1]), 0)  # restarts on page 2
            self.assertEqual(sorted(by_page[1]), list(range(len(by_page[1]))))

    def test_blank_line_does_not_inherit_neighbor_font_size(self):
        """Empty typeable lines must not be flagged using another paragraph’s 12pt runs."""
        parsed = {
            "text": "TITLE LINE ONE\n\nA Research Proposal",
            "metadata": {"format": "docx"},
            "paragraphs": [
                {
                    "paragraph_index": 0,
                    "source": "body",
                    "style": "Normal",
                    "text": "TITLE LINE ONE",
                    "lines": [{"line_index": 0, "text": "TITLE LINE ONE"}],
                    "runs": [{"text": "TITLE LINE ONE", "font": "Times New Roman", "size": 12, "bold": True}],
                    "formatting": {
                        "line_spacing": 1.5,
                        "alignment": "center",
                        "space_after_points": 36,
                    },
                },
                {
                    "paragraph_index": 1,
                    "source": "body",
                    "style": "Normal",
                    "text": "",
                    "lines": [{"line_index": 1, "text": ""}],
                    "runs": [{"text": "", "font": "Times New Roman", "size": 11}],
                    "formatting": {"line_spacing": 1.5, "alignment": "center"},
                },
                {
                    "paragraph_index": 2,
                    "source": "body",
                    "style": "Normal",
                    "text": "A Research Proposal",
                    "lines": [{"line_index": 2, "text": "A Research Proposal"}],
                    "runs": [{"text": "A Research Proposal", "font": "Times New Roman", "size": 11}],
                    "formatting": {"line_spacing": 1.5, "alignment": "center"},
                },
            ],
            "sections": [{
                "section_index": 0,
                "page_width_inches": 8.5,
                "page_height_inches": 11.0,
                "top_margin_inches": 1.0,
                "bottom_margin_inches": 1.0,
                "left_margin_inches": 1.0,
                "right_margin_inches": 1.0,
            }],
        }
        result = run_compliance_scan(
            parsed,
            {
                "font": {
                    "families": ["Times New Roman"],
                    "heading1_size": 12,
                    "heading2_size": 12,
                    "heading3_content_size": 11,
                    "sizes_points": [12, 11],
                }
            },
            "free",
        )
        # No font_size issue should claim an empty line measured 12pt.
        for issue in result["issues"]:
            if not str(issue["issue_type"]).startswith("font_size"):
                continue
            explanation = issue["explanation"].replace("12.0", "12")
            self.assertNotIn("Measured 12", explanation)
            self.assertNotIn("measured font size is 12", explanation.lower())

    def test_correct_docx_margins_score_100(self):
        parsed = {
            "text": "Body paragraph with enough words for scanning.",
            "metadata": {"format": "docx"},
            "paragraphs": [{
                "paragraph_index": 0,
                "source": "body",
                "style": "Normal",
                "text": "Body paragraph with enough words for scanning.",
                "lines": [{"line_index": 0, "text": "Body paragraph with enough words for scanning."}],
                "runs": [{"text": "Body paragraph with enough words for scanning.", "font": "Times New Roman", "size": 11}],
                "formatting": {"line_spacing": 1.5, "alignment": "left"},
            }],
            "sections": [{
                "section_index": 0,
                "page_width_inches": 8.5,
                "page_height_inches": 11.0,
                "top_margin_inches": 1.0,
                "bottom_margin_inches": 1.0,
                "left_margin_inches": 1.0,
                "right_margin_inches": 1.0,
            }],
        }
        result = run_compliance_scan(
            parsed,
            {
                "font": {"families": ["Times New Roman"], "sizes_points": [11]},
                "margins_inches": {"top": 1, "bottom": 1, "left": 1, "right": 1},
                "paper_size": {"name": "LETTER", "width_inches": 8.5, "height_inches": 11.0},
            },
            "free",
        )
        scores = {item["section"]: item["formatting_score"] for item in result["sections"]}
        self.assertEqual(scores["Margins"], 100.0)
        self.assertFalse([i for i in result["issues"] if "margin" in i["issue_type"]])

    def test_one_wrong_margin_does_not_zero_score(self):
        parsed = {
            "text": "Body paragraph with enough words for scanning.",
            "metadata": {"format": "docx"},
            "paragraphs": [{
                "paragraph_index": 0,
                "source": "body",
                "style": "Normal",
                "text": "Body paragraph with enough words for scanning.",
                "lines": [{"line_index": 0, "text": "Body paragraph with enough words for scanning."}],
                "runs": [{"text": "Body paragraph with enough words for scanning.", "font": "Times New Roman", "size": 11}],
                "formatting": {"line_spacing": 1.5, "alignment": "left"},
            }],
            "sections": [{
                "section_index": 0,
                "page_width_inches": 8.5,
                "page_height_inches": 11.0,
                "top_margin_inches": 1.0,
                "bottom_margin_inches": 1.0,
                "left_margin_inches": 0.5,
                "right_margin_inches": 1.0,
            }],
        }
        result = run_compliance_scan(
            parsed,
            {"margins_inches": {"top": 1, "bottom": 1, "left": 1, "right": 1}},
            "free",
        )
        scores = {item["section"]: item["formatting_score"] for item in result["sections"]}
        self.assertEqual(scores["Margins"], 75.0)
        left = [i for i in result["issues"] if i["issue_type"] == "left_margin"]
        self.assertEqual(len(left), 1)
        self.assertEqual(left[0]["severity"], "moderate")
        self.assertIn("0.5", left[0]["explanation"])

    def test_paper_size_mismatch_does_not_zero_correct_margins(self):
        parsed = {
            "text": "Body text here.",
            "metadata": {"format": "docx"},
            "paragraphs": [{
                "paragraph_index": 0,
                "source": "body",
                "text": "Body text here with enough words.",
                "lines": [{"line_index": 0, "text": "Body text here with enough words."}],
                "runs": [{"text": "Body text here with enough words.", "font": "Times New Roman", "size": 11}],
                "formatting": {},
            }],
            "sections": [{
                "section_index": 0,
                "page_width_inches": 8.27,
                "page_height_inches": 11.69,
                "top_margin_inches": 1.0,
                "bottom_margin_inches": 1.0,
                "left_margin_inches": 1.0,
                "right_margin_inches": 1.0,
            }],
        }
        result = run_compliance_scan(
            parsed,
            {
                "margins_inches": {"top": 1, "bottom": 1, "left": 1, "right": 1},
                "paper_size": {"name": "LETTER", "width_inches": 8.5, "height_inches": 11.0},
            },
            "free",
        )
        scores = {item["section"]: item["formatting_score"] for item in result["sections"]}
        self.assertEqual(scores["Margins"], 100.0)
        self.assertTrue([i for i in result["issues"] if i["issue_type"] == "paper_size"])

    def test_pdf_page_number_does_not_fail_bottom_margin(self):
        parsed = {
            "text": "Body line\n42",
            "metadata": {"format": "pdf"},
            "pages": [{
                "page_index": 0,
                "width_points": 612,
                "height_points": 792,
                "lines": [
                    {
                        "line_index": 0,
                        "text": "This body line sits on a normal academic page with enough words.",
                        "spans": [{"text": "body", "font": "TimesNewRoman", "size": 11}],
                        "bbox": [72, 72, 540, 86],
                    },
                    {
                        "line_index": 1,
                        "text": "42",
                        "spans": [{"text": "42", "font": "TimesNewRoman", "size": 11}],
                        "bbox": [300, 760, 320, 774],
                    },
                ],
            }],
        }
        result = run_compliance_scan(
            parsed,
            {"margins_inches": {"top": 1, "bottom": 1, "left": 1, "right": 1}},
            "free",
        )
        scores = {item["section"]: item["formatting_score"] for item in result["sections"]}
        self.assertEqual(scores["Margins"], 100.0)

    def test_docx_tables_headers_and_all_pages_are_scanned(self):
        from io import BytesIO

        from docx import Document
        from docx.shared import Pt

        document = Document()
        document.add_paragraph("Opening body paragraph with enough words for the scan.")
        table = document.add_table(rows=1, cols=1)
        run = table.cell(0, 0).paragraphs[0].add_run("Table cell uses Arial throughout this cell.")
        run.font.name = "Arial"
        run.font.size = Pt(11)
        nested = table.cell(0, 0).add_table(rows=1, cols=1)
        nested_run = nested.cell(0, 0).paragraphs[0].add_run("Nested table cell also uses Arial text.")
        nested_run.font.name = "Arial"
        nested_run.font.size = Pt(11)
        document.sections[0].header.is_linked_to_previous = False
        header_run = document.sections[0].header.paragraphs[0].add_run("University header running title")
        header_run.font.name = "Times New Roman"
        document.sections[0].footer.is_linked_to_previous = False
        footer_para = document.sections[0].footer.paragraphs[0]
        for run in list(footer_para.runs):
            run.text = ""
        footer_run = footer_para.add_run("Footer running text line")
        footer_run.font.name = "Times New Roman"
        buffer = BytesIO()
        document.save(buffer)
        parsed = parse_document(buffer.getvalue(), "docx")
        sources = {paragraph["source"] for paragraph in parsed["paragraphs"]}
        self.assertIn("body", sources)
        self.assertIn("table", sources)
        self.assertIn("header", sources)
        self.assertIn("footer", sources)
        combined = " ".join(paragraph["text"] for paragraph in parsed["paragraphs"])
        self.assertIn("Table cell uses Arial", combined)
        self.assertIn("Nested table cell", combined)
        self.assertIn("University header", combined)

        result = run_compliance_scan(
            parsed,
            {"font": {"families": ["Times New Roman"], "sizes_points": [11]}},
            "free",
        )
        family_issues = [issue for issue in result["issues"] if issue["issue_type"] == "font_family"]
        self.assertTrue(family_issues)
        self.assertGreaterEqual(family_issues[0]["count"], 2)

    def test_scan_keeps_locations_across_the_whole_document(self):
        paragraphs = []
        for index in range(150):
            paragraphs.append({
                "paragraph_index": index,
                "source": "body",
                "style": "Normal",
                "text": f"Body paragraph number {index + 1} uses Arial on purpose.",
                "lines": [{"line_index": index, "text": f"Body paragraph number {index + 1} uses Arial on purpose."}],
                "runs": [{"text": f"Body paragraph number {index + 1} uses Arial on purpose.", "font": "Arial", "size": 11}],
                "formatting": {"line_spacing": 1.5, "alignment": "left"},
            })
        parsed = {
            "text": "\n".join(item["text"] for item in paragraphs),
            "metadata": {"format": "docx"},
            "paragraphs": paragraphs,
            "sections": [{
                "section_index": 0,
                "page_width_inches": 8.5,
                "page_height_inches": 11.0,
                "top_margin_inches": 1.0,
                "bottom_margin_inches": 1.0,
                "left_margin_inches": 1.0,
                "right_margin_inches": 1.0,
            }],
        }
        result = run_compliance_scan(
            parsed,
            {"font": {"families": ["Times New Roman"], "sizes_points": [11]}},
            "free",
        )
        locations = result["issues"][0]["locations"]
        self.assertGreaterEqual(len(locations), 150)
        self.assertGreaterEqual(result["issues"][0]["count"], 150)
        self.assertGreater(result["page_count"], 1)
        last_page = max(loc["page"] for loc in locations)
        self.assertEqual(last_page, result["page_count"])


def _mechanics_paragraph(index, text, *, source="body", style="Normal", size=12, bold=False, italic=False, page_breaks=0, drawing=False):
    return {
        "paragraph_index": index,
        "source": source,
        "style": style,
        "text": text,
        "has_drawing": drawing,
        "lines": [{"line_index": 0, "text": text}],
        "runs": [{
            "text": text,
            "font": "Times New Roman",
            "size": size,
            "bold": bold,
            "italic": italic,
            "rendered_page_breaks": page_breaks,
        }],
        "formatting": {"line_spacing": 2.0, "alignment": "left"},
    }


class LayoutAndPageTests(unittest.TestCase):
    def _doc(self, paragraphs):
        return {
            "text": "\n".join(p["text"] for p in paragraphs),
            "metadata": {"format": "docx"},
            "paragraphs": paragraphs,
            "sections": [{
                "section_index": 0,
                "page_width_inches": 8.5,
                "page_height_inches": 11.0,
                "top_margin_inches": 1.0,
                "bottom_margin_inches": 1.0,
                "left_margin_inches": 1.0,
                "right_margin_inches": 1.0,
                "start_type": "new_page",
            }],
        }

    def test_table_and_figure_rules_flag_real_violations(self):
        paragraphs = [
            _mechanics_paragraph(0, "Table 1. Sensor readings"),
            _mechanics_paragraph(1, "Year Value", source="table", size=10),
            _mechanics_paragraph(2, "Fig. 1. Setup"),
            _mechanics_paragraph(3, "", drawing=True),
            _mechanics_paragraph(
                4,
                "Table 2 shows the sample in the narrative and is not a caption.",
            ),
        ]
        rules = {
            "font": {"families": ["Times New Roman"], "body_size": 12},
            "table_layout_requirements": [
                "Each table must be labeled Table number with its title in capital letters, both placed above the table."
            ],
            "figure_layout_requirements": [
                "Each figure must be labeled Figure number with its title placed below the figure."
            ],
        }
        issues = run_compliance_scan(self._doc(paragraphs), rules, "free")["issues"]
        types = {issue["issue_type"] for issue in issues}
        self.assertIn("table_caption", types)
        self.assertIn("figure_caption", types)
        self.assertIn("figure_caption_placement", types)
        prose = [
            issue for issue in issues
            if any("narrative" in (loc.get("excerpt") or "") for loc in issue.get("locations") or [])
        ]
        self.assertEqual(prose, [])

    def test_matching_table_and_figure_captions_are_not_flagged(self):
        paragraphs = [
            _mechanics_paragraph(0, "Table 1. SENSOR READINGS"),
            _mechanics_paragraph(1, "Year Value", source="table", size=12),
            _mechanics_paragraph(2, "", drawing=True),
            _mechanics_paragraph(3, "Figure 1. Setup"),
        ]
        rules = {
            "font": {"families": ["Times New Roman"], "body_size": 12, "sizes_points": [12]},
            "table_layout_requirements": [
                "Each table must be labeled Table number with its title in capital letters, both placed above the table."
            ],
            "figure_layout_requirements": [
                "Each figure must be labeled Figure number with its title placed below the figure."
            ],
        }
        types = {issue["issue_type"] for issue in run_compliance_scan(self._doc(paragraphs), rules, "free")["issues"]}
        self.assertFalse(any(item.startswith("table_caption") or item.startswith("figure_caption") for item in types))

    def test_body_page_is_not_labeled_as_the_title_page(self):
        prose = (
            "This paragraph explains the method in enough detail to count as ordinary "
            "body text on a normal manuscript page."
        )
        paragraphs = [
            _mechanics_paragraph(0, "A Research Title", size=16, bold=True),
            _mechanics_paragraph(1, "Chapter 1", style="Heading 1", size=14, bold=True, page_breaks=1),
            _mechanics_paragraph(2, prose, size=11),
            _mechanics_paragraph(3, "Year Value", source="table", size=10),
        ]
        rules = {
            "font": {
                "families": ["Times New Roman"],
                "body_size": 12,
                "heading1_size": 14,
                "heading_levels": 1,
                "sizes_points": [12, 14],
            },
            "pagination": {"title_page": "Number shown"},
        }
        issues = run_compliance_scan(self._doc(paragraphs), rules, "free")["issues"]
        for issue in issues:
            if "title page" not in issue["title"].lower():
                continue
            pages = {loc.get("page") for loc in issue.get("locations") or []}
            self.assertEqual(pages, {1})
        body_font = [issue for issue in issues if issue["title"] == "body text font size differs"]
        self.assertTrue(body_font)
        body_pages = {loc.get("page") for issue in body_font for loc in issue["locations"]}
        self.assertIn(2, body_pages)
        self.assertNotIn(1, body_pages)

    def test_heading_font_issues_use_heading_label_and_specific_size(self):
        paragraphs = [
            _mechanics_paragraph(0, "A Research Title", size=16, bold=True),
            _mechanics_paragraph(1, "CHAPTER 1", style="Heading 1", size=12, bold=True, page_breaks=1),
            _mechanics_paragraph(2, "Method", style="Heading 2", size=11, bold=True),
            _mechanics_paragraph(
                3,
                "This paragraph explains the method in enough detail to count as ordinary body text.",
                size=11,
            ),
        ]
        rules = {
            "font": {
                "families": ["Times New Roman"],
                "body_size": 12,
                "heading1_size": 14,
                "heading2_size": 12,
                "heading_levels": 2,
                "sizes_points": [11, 12, 14],
            },
            "pagination": {"title_page": "Number shown"},
        }
        issues = run_compliance_scan(self._doc(paragraphs), rules, "free")["issues"]
        by_type = {issue["issue_type"]: issue for issue in issues}
        self.assertIn("font_size_heading1", by_type)
        self.assertIn("font_size_heading2", by_type)
        h1 = by_type["font_size_heading1"]
        h2 = by_type["font_size_heading2"]
        self.assertEqual(h1["title"], "Heading 1 font size differs")
        self.assertEqual(h2["title"], "Heading 2 font size differs")
        self.assertNotIn("title page", h1["title"].lower())
        self.assertNotIn("title page", h2["title"].lower())
        self.assertIn("14 pt", h1["explanation"])
        self.assertIn("12 pt", h2["explanation"])
        self.assertIn("measured font size", h1["explanation"].lower())
        self.assertNotIn("11", h1["explanation"].split("require", 1)[-1])
        self.assertNotIn("14", h2["explanation"].split("require", 1)[-1])
        for loc in h1["locations"]:
            self.assertNotIn("title page", (loc.get("excerpt") or "").lower())
        excerpts = " ".join(loc.get("excerpt") or "" for loc in h1["locations"])
        self.assertIn("CHAPTER 1", excerpts)

    def test_body_size_does_not_flag_heading_styles_when_heading_sizes_unset(self):
        paragraphs = [
            _mechanics_paragraph(0, "CHAPTER 1 INTRODUCTION", style="Heading 1", size=16, bold=True),
            _mechanics_paragraph(1, "Related Work", style="Heading 2", size=14, bold=True),
            _mechanics_paragraph(
                2,
                "This paragraph explains the method in enough detail to count as ordinary body text.",
                size=12,
            ),
        ]
        rules = normalize_mechanics_rules({"font": {"families": ["Times New Roman"], "body_size": 12}})
        issues = run_compliance_scan(self._doc(paragraphs), rules, "free")["issues"]
        titles = {issue["title"] for issue in issues}
        self.assertFalse(any("Heading" in title and "font size" in title for title in titles))
        self.assertNotIn("body text font size differs", titles)

    def test_sizes_points_alone_does_not_flag_title_when_headings_unset(self):
        """No heading rules + sizes_points body fallback must not flag a larger title line."""
        paragraphs = [
            _mechanics_paragraph(
                0,
                "PAPERPILOT: AN AI-POWERED ACADEMIC DOCUMENT ANALYSIS AND COMPLIANCE CHECKING SYSTEM",
                size=14,
                bold=True,
            ),
            _mechanics_paragraph(
                1,
                "This paragraph explains the method in enough detail to count as ordinary body text.",
                size=11,
            ),
        ]
        rules = normalize_mechanics_rules(
            {"font": {"families": ["Arial"], "sizes_points": [11.0], "color": "Automatic"}}
        )
        self.assertNotIn("heading1_size", rules.get("font") or {})
        issues = run_compliance_scan(self._doc(paragraphs), rules, "free")["issues"]
        types = {issue["issue_type"] for issue in issues}
        titles = {issue["title"] for issue in issues}
        self.assertNotIn("font_size_heading1", types)
        self.assertFalse(any("Heading" in title and "font size" in title for title in titles))
        # Title must not be reported as a body-size failure either.
        body_issues = [issue for issue in issues if issue.get("issue_type") == "font_size"]
        for issue in body_issues:
            excerpts = " ".join(loc.get("excerpt") or "" for loc in issue.get("locations") or [])
            self.assertNotIn("PAPERPILOT", excerpts.upper())

    def test_empty_heading1_slot_is_not_flagged_by_sizes_points(self):
        """UI may show an empty Heading 1 card when only Heading 2 has rules.
        sizes_points must not invent an H1 size requirement."""
        paragraphs = [
            _mechanics_paragraph(
                0,
                "PAPERPILOT: AN AI-POWERED ACADEMIC DOCUMENT",
                style="Heading 1",
                size=14,
                bold=True,
            ),
            _mechanics_paragraph(1, "Methods", style="Heading 2", size=12, bold=True),
            _mechanics_paragraph(
                2,
                "This paragraph explains the method in enough detail to count as ordinary body text.",
                size=11,
            ),
        ]
        rules = normalize_mechanics_rules({
            "font": {
                "families": ["Times New Roman"],
                "heading2_size": 12,
                "heading_levels": 2,
                "sizes_points": [11.0, 12.0],
                "color": "Black/Automatic",
            }
        })
        issues = run_compliance_scan(self._doc(paragraphs), rules, "free")["issues"]
        types = {issue["issue_type"] for issue in issues}
        titles = {issue["title"] for issue in issues}
        self.assertNotIn("font_size_heading1", types)
        self.assertFalse(any("Heading 1" in title and "font size" in title for title in titles))
        # sizes_points alone also must not label unset H1 as wrong
        only_listed = normalize_mechanics_rules({
            "font": {
                "families": ["Times New Roman"],
                "heading_levels": 2,
                "sizes_points": [11.0, 12.0],
            }
        })
        listed_issues = run_compliance_scan(self._doc(paragraphs), only_listed, "free")["issues"]
        listed_types = {issue["issue_type"] for issue in listed_issues}
        self.assertNotIn("font_size_heading1", listed_types)
        self.assertFalse(
            any("Heading 1" in issue["title"] and "font size" in issue["title"] for issue in listed_issues)
        )

    def test_heading2_and_heading3_sizes_do_not_flag_unset_heading1(self):
        """Only H2=11 and H3=12 set — a 14pt Heading 1 must pass (no H1 size rule)."""
        paragraphs = [
            _mechanics_paragraph(
                0,
                "PAPERPILOT: AN AI-POWERED ACADEMIC DOCUMENT",
                style="Heading 1",
                size=14,
                bold=True,
            ),
            _mechanics_paragraph(1, "Methods", style="Heading 2", size=11, bold=True),
            _mechanics_paragraph(2, "Instruments", style="Heading 3", size=12, bold=True),
            _mechanics_paragraph(
                3,
                "This paragraph explains the method in enough detail to count as ordinary body text.",
                size=12,
            ),
        ]
        rules = normalize_mechanics_rules({
            "font": {
                "families": ["Times New Roman"],
                "heading2_size": 11,
                "heading3_size": 12,
                "heading_levels": 3,
                # Stale client payloads used to dump heading sizes into sizes_points.
                "sizes_points": [11.0, 12.0],
            }
        })
        self.assertNotIn("sizes_points", rules.get("font") or {})
        self.assertNotIn("heading1_size", rules.get("font") or {})
        issues = run_compliance_scan(self._doc(paragraphs), rules, "free")["issues"]
        self.assertNotIn("font_size_heading1", {issue["issue_type"] for issue in issues})
        self.assertFalse(
            any("Heading 1" in (issue.get("title") or "") and "font size" in (issue.get("title") or "")
                for issue in issues)
        )

    def test_heading2_mismatch_still_flags_when_heading1_unset(self):
        """H1 unset passes; wrong H2 size still appears on Issues."""
        paragraphs = [
            _mechanics_paragraph(
                0,
                "PAPERPILOT: AN AI-POWERED ACADEMIC DOCUMENT",
                style="Heading 1",
                size=14,
                bold=True,
            ),
            _mechanics_paragraph(1, "Methods", style="Heading 2", size=14, bold=True),
            _mechanics_paragraph(2, "Instruments", style="Heading 3", size=12, bold=True),
            _mechanics_paragraph(
                3,
                "This paragraph explains the method in enough detail to count as ordinary body text.",
                size=12,
            ),
        ]
        rules = normalize_mechanics_rules({
            "font": {
                "families": ["Times New Roman"],
                "heading2_size": 11,
                "heading3_size": 12,
                "heading_levels": 3,
                "sizes_points": [11.0, 12.0],
            }
        })
        issues = run_compliance_scan(self._doc(paragraphs), rules, "free")["issues"]
        types = {issue["issue_type"] for issue in issues}
        self.assertNotIn("font_size_heading1", types)
        self.assertIn("font_size_heading2", types)
        h2 = next(issue for issue in issues if issue["issue_type"] == "font_size_heading2")
        self.assertIn("Heading 2", h2["title"])
        self.assertNotIn("Heading 1", h2["title"])

    def test_heading1_rule_does_not_flag_heading2_or_heading3(self):
        paragraphs = [
            _mechanics_paragraph(0, "CHAPTER 1", style="Heading 1", size=12, bold=True),
            _mechanics_paragraph(1, "Methods", style="Heading 2", size=12, bold=True),
            _mechanics_paragraph(2, "Instruments", style="Heading 3", size=12, bold=True),
            _mechanics_paragraph(
                3,
                "This paragraph explains the method in enough detail to count as ordinary body text.",
                size=12,
            ),
        ]
        rules = normalize_mechanics_rules({
            "font": {
                "families": ["Times New Roman"],
                "heading1_size": 14,
                "heading2_size": 12,
                "heading3_size": 12,
                "body_size": 12,
                "heading_levels": 3,
            }
        })
        issues = run_compliance_scan(self._doc(paragraphs), rules, "free")["issues"]
        by_type = {issue["issue_type"]: issue for issue in issues}
        self.assertIn("font_size_heading1", by_type)
        self.assertNotIn("font_size_heading2", by_type)
        self.assertNotIn("font_size_heading3", by_type)
        excerpts = " ".join(
            loc.get("excerpt") or ""
            for loc in by_type["font_size_heading1"].get("locations") or []
        )
        self.assertIn("CHAPTER 1", excerpts)
        self.assertNotIn("Methods", excerpts)
        self.assertNotIn("Instruments", excerpts)

    def test_acknowledgement_body_lines_are_not_heading1(self):
        """Title + multi-line indented body: only ACKNOWLEDGEMENT is a heading."""
        body = (
            "I would like to express my sincere gratitude to the panel members whose insightful "
            "suggestions and helpful comments contributed greatly to the refinement and quality of this study."
        )
        paragraphs = [
            _mechanics_paragraph(0, "ACKNOWLEDGEMENT", style="Heading 1", size=14, bold=True),
            _mechanics_paragraph(1, body, size=14, bold=True),
        ]
        paragraphs[1]["formatting"] = {
            "line_spacing": 2.0,
            "alignment": "justify",
            "first_line_indent_inches": 0.5,
        }
        # Soft-wrap the bold body under Heading 1 style as well (false style on prose).
        long_h1 = (body + " ") * 2
        paragraphs.append(
            _mechanics_paragraph(2, long_h1, style="Heading 1", size=14, bold=False)
        )
        paragraphs[2]["formatting"] = {
            "line_spacing": 2.0,
            "alignment": "justify",
            "first_line_indent_inches": 0.5,
        }
        rules = normalize_mechanics_rules(
            {
                "font": {
                    "families": ["Times New Roman"],
                    "heading1_size": 30,
                    "body_size": 14,
                    "heading_levels": 1,
                }
            }
        )
        result = run_compliance_scan(self._doc(paragraphs), rules, "free")
        heading_issues = [
            issue
            for issue in result["issues"]
            if issue.get("issue_type") == "font_size_heading1"
        ]
        self.assertEqual(len(heading_issues), 1)
        excerpts = " ".join(
            loc.get("excerpt") or "" for loc in heading_issues[0].get("locations") or []
        )
        self.assertIn("ACKNOWLEDGEMENT", excerpts)
        self.assertNotIn("helpful comments", excerpts)
        self.assertIn("30", heading_issues[0]["explanation"])

    def test_real_heading1_wrong_size_is_one_issue_on_that_line(self):
        paragraphs = [
            _mechanics_paragraph(0, "CHAPTER 1", style="Heading 1", size=12, bold=True),
            _mechanics_paragraph(
                1,
                "This paragraph explains the method in enough detail to count as ordinary body text.",
                size=12,
            ),
        ]
        rules = normalize_mechanics_rules(
            {
                "font": {
                    "families": ["Times New Roman"],
                    "heading1_size": 16,
                    "body_size": 12,
                    "heading_levels": 1,
                }
            }
        )
        issues = run_compliance_scan(self._doc(paragraphs), rules, "free")["issues"]
        h1 = [issue for issue in issues if issue.get("issue_type") == "font_size_heading1"]
        self.assertEqual(len(h1), 1)
        self.assertEqual(len(h1[0]["locations"]), 1)
        self.assertIn("16 pt", h1[0]["explanation"])
        self.assertIn("CHAPTER 1", h1[0]["locations"][0].get("excerpt") or "")

    def test_numbered_and_chapter_titles_still_detected(self):
        paragraphs = [
            _mechanics_paragraph(0, "1.1 Background of the Study", size=12, bold=True),
            _mechanics_paragraph(1, "CHAPTER 2", size=12, bold=True),
            _mechanics_paragraph(2, "REFERENCES", size=12, bold=True),
            _mechanics_paragraph(
                3,
                "This paragraph explains the method in enough detail to count as ordinary body text.",
                size=12,
            ),
        ]
        rules = normalize_mechanics_rules(
            {
                "font": {
                    "families": ["Times New Roman"],
                    "heading1_size": 14,
                    "body_size": 12,
                    "heading_levels": 1,
                }
            }
        )
        from app.compliance import _annotate_roles, _document_units

        units = _document_units(self._doc(paragraphs), rules)
        _annotate_roles(units, rules)
        roles = {
            (unit.get("text") or "").strip(): unit.get("role")
            for unit in units
            if (unit.get("text") or "").strip()
        }
        self.assertEqual(roles.get("1.1 Background of the Study"), "heading1")
        self.assertEqual(roles.get("CHAPTER 2"), "heading1")
        self.assertEqual(roles.get("REFERENCES"), "heading1")

    def test_multiline_all_caps_prose_is_not_a_heading(self):
        prose = (
            "THE COMPLETION OF THIS RESEARCH STUDY WOULD NOT HAVE BEEN POSSIBLE WITHOUT THE "
            "GUIDANCE AND SUPPORT OF MANY PEOPLE WHO CONTRIBUTED HELPFUL COMMENTS."
        )
        paragraphs = [
            _mechanics_paragraph(0, prose, size=12, bold=True),
            _mechanics_paragraph(
                1,
                "This paragraph explains the method in enough detail to count as ordinary body text.",
                size=12,
            ),
        ]
        paragraphs[0]["formatting"] = {
            "line_spacing": 2.0,
            "alignment": "justify",
            "first_line_indent_inches": 0.5,
        }
        rules = normalize_mechanics_rules(
            {
                "font": {
                    "families": ["Times New Roman"],
                    "heading1_size": 30,
                    "body_size": 12,
                    "heading_levels": 1,
                }
            }
        )
        issues = run_compliance_scan(self._doc(paragraphs), rules, "free")["issues"]
        self.assertFalse(
            any(issue.get("issue_type") == "font_size_heading1" for issue in issues)
        )

    def test_no_heading1_size_means_zero_heading_size_issues(self):
        paragraphs = [
            _mechanics_paragraph(0, "ACKNOWLEDGEMENT", style="Heading 1", size=14, bold=True),
            _mechanics_paragraph(
                1,
                "This paragraph explains the method in enough detail to count as ordinary body text.",
                size=12,
            ),
        ]
        rules = normalize_mechanics_rules(
            {"font": {"families": ["Times New Roman"], "body_size": 12, "heading_levels": 1}}
        )
        result = run_compliance_scan(self._doc(paragraphs), rules, "free")
        issues = result["issues"]
        self.assertFalse(
            any(
                "Heading" in (issue.get("title") or "") and "font size" in (issue.get("title") or "")
                for issue in issues
            )
        )
        sections = {item["section"]: item for item in result.get("sections") or []}
        fonts = sections.get("Fonts") or {}
        if fonts.get("units_checked"):
            self.assertGreaterEqual(fonts["units_checked"], fonts.get("units_failed") or 0)

    def test_arialmt_reported_as_arial_with_detected_font_wording(self):
        paragraphs = [
            _mechanics_paragraph(0, "This paragraph uses a PDF PostScript Arial face on purpose.", size=12),
        ]
        paragraphs[0]["runs"][0]["font"] = "ArialMT"
        rules = normalize_mechanics_rules({"font": {"families": ["Times New Roman"], "body_size": 12}})
        issues = run_compliance_scan(self._doc(paragraphs), rules, "free")["issues"]
        family = next(issue for issue in issues if issue["issue_type"] == "font_family")
        self.assertIn("detected font is Arial", family["explanation"])
        self.assertNotIn("ArialMT", family["explanation"])
        self.assertNotIn("Measured font", family["explanation"])

    def test_string_table_layout_rule_still_flags_caption_case(self):
        paragraphs = [
            _mechanics_paragraph(0, "Table 1. Sensor readings"),
            _mechanics_paragraph(1, "Year Value", source="table", size=10),
        ]
        rules = {
            "font": {"families": ["Times New Roman"], "body_size": 12},
            # Legacy / client payloads sometimes store a plain string instead of a list.
            "table_layout_requirements": (
                "Each table must be labeled Table number with its title in capital letters, "
                "both placed above the table."
            ),
        }
        types = {issue["issue_type"] for issue in run_compliance_scan(self._doc(paragraphs), rules, "free")["issues"]}
        self.assertIn("table_caption", types)

    def test_heading_only_rules_do_not_invent_body_size(self):
        rules = normalize_mechanics_rules({
            "font": {
                "families": ["Times New Roman"],
                "heading1_size": 19,
                "heading2_size": 14,
                "heading3_size": 12,
                "sizes_points": [19, 14, 12],
                "heading3_content_size": 19,
            }
        })
        font = rules.get("font") or {}
        self.assertNotIn("body_size", font)
        self.assertNotIn("heading3_content_size", font)

    def test_list_valued_heading1_size_is_unset_and_never_prints_brackets(self):
        """Stale payloads may store heading1_size as [11.0]; that must not create an issue."""
        paragraphs = [
            _mechanics_paragraph(
                0,
                "PAPERPILOT: AN AI-POWERED ACADEMIC DOCUMENT ANALYSIS AND COMPLIANCE CHECKING SYSTEM",
                style="Heading 1",
                size=14,
                bold=True,
            ),
            _mechanics_paragraph(
                1,
                "This paragraph explains the method in enough detail to count as ordinary body text.",
                size=11,
            ),
        ]
        # Bypass normalize so the list shape reaches the checker (legacy / hand-edited JSON).
        rules = {
            "font": {
                "families": ["Arial"],
                "heading1_size": [11.0],
                "sizes_points": [11.0],
                "color": "Black/Automatic",
            }
        }
        from app.compliance import (
            _filter_unset_heading_font_issues,
            _format_expected_pt,
            _heading_sizes,
            _scalar_pt,
        )

        self.assertEqual(_heading_sizes(rules["font"]), {})
        self.assertIsNone(_scalar_pt([11.0]))
        # Issue text may join list items as scalars — never as "[11.0]".
        self.assertEqual(_format_expected_pt([11.0]), "11")
        self.assertNotIn("[", _format_expected_pt([11.0]) or "")
        result = run_compliance_scan(self._doc(paragraphs), rules, "free")
        types = {issue["issue_type"] for issue in result["issues"]}
        self.assertNotIn("font_size_heading1", types)
        for issue in result["issues"]:
            blob = " ".join(
                str(issue.get(key) or "")
                for key in ("title", "explanation", "recommendation", "summary")
            )
            self.assertNotIn("[11.0]", blob)
            self.assertNotIn("[11]", blob)
        # Enrichment-style list leak is stripped by the post-filter.
        leaked = [
            {
                "issue_type": "font_size",
                "title": "body text font size differs",
                "explanation": "Measured 14 pt; the format specifies [11.0] pt for Heading 1.",
                "recommendation": "Set to [11.0] pt.",
                "summary": "Heading 1 size wrong",
                "locations": [],
            }
        ]
        cleaned = _filter_unset_heading_font_issues(leaked, rules)
        self.assertEqual(len(cleaned), 1)
        self.assertNotIn("[11.0]", cleaned[0]["explanation"])
        self.assertNotIn("Heading 1", cleaned[0]["explanation"])

    def test_ignore_heading1_size_counts_as_pass(self):
        paragraphs = [
            _mechanics_paragraph(0, "CHAPTER 1", style="Heading 1", size=14, bold=True),
            _mechanics_paragraph(
                1,
                "This paragraph explains the method in enough detail to count as ordinary body text.",
                size=12,
            ),
        ]
        rules = {
            "font": {
                "families": ["Times New Roman"],
                "heading1_size": "IGNORE",
                "body_size": 12,
            }
        }
        result = run_compliance_scan(self._doc(paragraphs), rules, "free")
        self.assertNotIn(
            "font_size_heading1",
            {issue["issue_type"] for issue in result["issues"]},
        )
        fonts = next(
            item for item in result["scoring"]["categories"] if item["category"] == "Fonts"
        )
        self.assertGreater(fonts["units_checked"], 0)
        self.assertEqual(fonts["units_checked"], fonts["units_passed"] + fonts["units_failed"])

    def test_heading1_set_to_16_wrong_size_says_16_pt(self):
        paragraphs = [
            _mechanics_paragraph(0, "CHAPTER 1", style="Heading 1", size=12, bold=True),
            _mechanics_paragraph(
                1,
                "This paragraph explains the method in enough detail to count as ordinary body text.",
                size=12,
            ),
        ]
        rules = normalize_mechanics_rules(
            {"font": {"families": ["Times New Roman"], "heading1_size": 16, "body_size": 12}}
        )
        issues = run_compliance_scan(self._doc(paragraphs), rules, "free")["issues"]
        h1 = [issue for issue in issues if issue["issue_type"] == "font_size_heading1"]
        self.assertEqual(len(h1), 1)
        self.assertIn("16 pt", h1[0]["explanation"])
        self.assertNotIn("[", h1[0]["explanation"])

    def test_black_automatic_color_is_enforced_not_unset(self):
        paragraphs = [
            _mechanics_paragraph(
                0,
                "This paragraph explains the method in enough detail to count as ordinary body text.",
                size=12,
            ),
        ]
        paragraphs[0]["runs"][0]["color"] = "0000FF"
        rules = normalize_mechanics_rules(
            {"font": {"families": ["Times New Roman"], "body_size": 12, "color": "Black/Automatic"}}
        )
        types = {issue["issue_type"] for issue in run_compliance_scan(self._doc(paragraphs), rules, "free")["issues"]}
        self.assertIn("font_color", types)

    def test_zero_indent_rule_is_enforced(self):
        paragraphs = [
            _mechanics_paragraph(
                0,
                "This paragraph explains the method in enough detail to count as ordinary body text "
                "and continues onto a second visual line for indentation measurement.",
                size=12,
            ),
            _mechanics_paragraph(
                1,
                "Second line of the same paragraph for layout.",
                size=12,
            ),
        ]
        # Same paragraph_index so they form one prose paragraph with 2+ lines.
        paragraphs[0]["paragraph_index"] = 0
        paragraphs[1]["paragraph_index"] = 0
        paragraphs[0]["formatting"] = {
            "first_line_indent_inches": 0.5,
            "left_indent_inches": 0.0,
            "alignment": "left",
            "line_spacing": 2.0,
        }
        paragraphs[1]["formatting"] = {
            "first_line_indent_inches": 0.0,
            "left_indent_inches": 0.0,
            "alignment": "left",
            "line_spacing": 2.0,
        }
        paragraphs[0]["bbox"] = [72, 72, 500, 86]
        paragraphs[1]["bbox"] = [72, 100, 500, 114]
        rules = normalize_mechanics_rules(
            {"font": {"families": ["Times New Roman"], "body_size": 12}, "first_line_indent_inches": 0}
        )
        self.assertEqual(rules.get("first_line_indent_inches"), 0.0)
        types = {issue["issue_type"] for issue in run_compliance_scan(self._doc(paragraphs), rules, "free")["issues"]}
        self.assertIn("first_line_indent", types)

    def test_ignore_line_spacing_skips_spacing_check(self):
        paragraphs = [
            _mechanics_paragraph(
                0,
                "This paragraph explains the method in enough detail to count as ordinary body text.",
                size=12,
            ),
        ]
        paragraphs[0]["formatting"] = {"line_spacing": 1.0, "alignment": "left"}
        result = run_compliance_scan(
            self._doc(paragraphs),
            {"font": {"body_size": 12}, "line_spacing": "IGNORE"},
            "free",
        )
        spacing = next(
            item for item in result["scoring"]["categories"] if item["category"] == "Spacing"
        )
        self.assertEqual(spacing["status"], "not_applicable")
        self.assertNotIn("line_spacing", {issue["issue_type"] for issue in result["issues"]})


class PersistenceTests(FirebaseTestCase):
    def parsed(self):
        return {
            "text": "Example manuscript text.",
            "metadata": {"format": "docx", "pagination_fidelity": "explicit_breaks_only_no_static_layout"},
            "paragraphs": [],
            "sections": [],
        }

    def mechanics(self, name="Guide"):
        parsed = self.parsed()
        return compliance_db.create_mechanics(
            "owner", name, "guide.docx", "docx", "Use Arial.", parsed,
            {"font": {"families": ["Arial"]}},
        )

    def test_relationships_and_transactional_scan_usage(self):
        parsed = self.parsed()
        mechanics = self.mechanics()
        version, created_parent = compliance_db.create_version(
            "owner", "Paper", mechanics["id"], "paper.docx", "docx",
            parsed["text"], parsed, None,
        )
        self.assertTrue(created_parent)
        with self.assertRaises(LookupError):
            compliance_db.persist_scan(
                "owner", version["manuscript_id"], "missing", mechanics["id"],
                100.0, [], [],
            )
        self.assertEqual(compliance_db.subscription_snapshot("owner")["used"], 0)
        scan = compliance_db.persist_scan(
            "owner", version["manuscript_id"], version["id"], mechanics["id"], 100.0, [],
            [{"section": "General", "formatting_score": 100.0, "issue_count": 0, "issues": []}],
        )
        self.assertEqual(scan["manuscript_version_id"], version["id"])
        self.assertEqual(compliance_db.subscription_snapshot("owner")["used"], 1)

    def test_mechanics_names_are_unique_and_rename_persists_filename(self):
        first = self.mechanics("Graduate Guide")
        with self.assertRaises(compliance_db.MechanicsNameConflict):
            self.mechanics("graduate guide")
        renamed = compliance_db.rename_mechanics("owner", first["id"], "Thesis Guide")
        self.assertEqual(renamed["source_filename"], "Thesis Guide.docx")
        self.assertEqual(
            compliance_db.list_mechanics("owner")[0]["source_filename"],
            "Thesis Guide.docx",
        )
        self.assertTrue(compliance_db.delete_mechanics("owner", first["id"]))

    def test_deleted_mechanics_name_can_be_reused(self):
        first = self.mechanics("Format Guide")
        self.assertTrue(compliance_db.delete_mechanics("owner", first["id"]))
        self.assertEqual(compliance_db.list_mechanics("owner"), [])
        # Same display name must be allowed again after delete.
        second = self.mechanics("Format Guide")
        self.assertNotEqual(first["id"], second["id"])
        self.assertEqual(second["name"], "Format Guide")

    def test_same_title_upload_becomes_new_version(self):
        parsed, mechanics = self.parsed(), self.mechanics()
        first, created_first = compliance_db.create_version(
            "owner", "Thesis Draft", mechanics["id"], "a.docx", "docx", "Text", parsed, None
        )
        self.assertTrue(created_first)
        self.assertTrue(first["pending"])
        compliance_db.persist_scan(
            "owner", first["manuscript_id"], first["id"], mechanics["id"], 90.0, [],
            [{"section": "General", "formatting_score": 90.0, "issue_count": 0, "issues": []}],
        )
        saved = compliance_db.get_version("owner", first["manuscript_id"], first["id"])
        self.assertFalse(saved["pending"])
        second, created_second = compliance_db.create_version(
            "owner", "thesis draft", mechanics["id"], "b.docx", "docx", "Text", parsed, None
        )
        self.assertFalse(created_second)
        self.assertEqual(first["manuscript_id"], second["manuscript_id"])
        self.assertEqual(first["version_number"], 1)
        self.assertEqual(second["version_number"], 2)
        listed = compliance_db.list_manuscripts("owner")
        self.assertEqual(listed[0]["version_count"], 1)
        self.assertEqual(listed[0]["current_version_number"], 1)

    def test_unscanned_upload_releases_its_version_number(self):
        parsed, mechanics = self.parsed(), self.mechanics()
        first, _ = compliance_db.create_version(
            "owner", "Thesis Draft", mechanics["id"], "a.docx", "docx", "Text", parsed, None
        )
        compliance_db.persist_scan(
            "owner", first["manuscript_id"], first["id"], mechanics["id"], 90.0, [],
            [{"section": "General", "formatting_score": 90.0, "issue_count": 0, "issues": []}],
        )
        abandoned, _ = compliance_db.create_version(
            "owner", "Thesis Draft", mechanics["id"], "b.docx", "docx", "Text", parsed, first["manuscript_id"]
        )
        self.assertEqual(abandoned["version_number"], 2)
        self.assertTrue(abandoned["pending"])
        self.assertEqual(compliance_db.list_manuscripts("owner")[0]["current_version_number"], 1)
        self.assertTrue(compliance_db.delete_version("owner", first["manuscript_id"], abandoned["id"]))
        self.assertEqual(compliance_db.list_manuscripts("owner")[0]["current_version_number"], 1)
        again, _ = compliance_db.create_version(
            "owner", "Thesis Draft", mechanics["id"], "c.docx", "docx", "Text", parsed, first["manuscript_id"]
        )
        self.assertEqual(again["version_number"], 2)
        self.assertTrue(again["pending"])

    def test_orphaned_mechanics_name_reservation_is_reclaimed(self):
        # Simulate a stale name index entry whose mechanics row is already gone.
        name_ref = compliance_db._reference(
            f"{compliance_db.ROOT}/mechanics_names/owner/{compliance_db._name_key('Orphan Guide')}"
        )
        name_ref.set("missing-mechanics-id")
        reclaimed = self.mechanics("Orphan Guide")
        self.assertEqual(reclaimed["name"], "Orphan Guide")
        self.assertEqual(name_ref.get(), reclaimed["id"])

    def test_mechanics_used_by_version_can_still_be_deleted(self):
        parsed, mechanics = self.parsed(), self.mechanics()
        version, _ = compliance_db.create_version(
            "owner", "Paper", mechanics["id"], "paper.docx", "docx", "Text", parsed, None
        )
        self.assertEqual(
            compliance_db.get_version("owner", version["manuscript_id"], version["id"])["mechanics_id"],
            mechanics["id"],
        )
        # Deletion is now always allowed regardless of whether the mechanics is in use.
        self.assertTrue(compliance_db.delete_mechanics("owner", mechanics["id"]))
        # And the name must be free for a fresh upload.
        again = self.mechanics()
        self.assertTrue(again["id"])

    def test_username_reservation_is_case_insensitive(self):
        self.assertTrue(profiles.reserve_username("Pilot.User", "one"))
        self.assertTrue(profiles.username_taken("pilot.user"))
        self.assertFalse(profiles.reserve_username("PILOT.USER", "two"))

    def _activate(self, cs="cs_1", pay="pay_1"):
        return compliance_db.activate_premium_from_payment(
            "owner", billing_period="monthly", payment_method="card", amount_pesos=949,
            checkout_session_id=cs, payment_id=pay,
        )

    def test_payment_is_applied_once(self):
        first = self._activate()
        renews = first["renews_at"]
        self.assertEqual(first["tier"], "premium")
        again = self._activate()
        self.assertEqual(again["renews_at"], renews)
        subscribed = [h for h in again["history"] if h.get("action") == "subscribed"]
        self.assertEqual(len(subscribed), 1)

    def test_concurrent_activation_claim_blocks_second_caller(self):
        self.assertIsNotNone(compliance_db._claim_payment_activation("cs_2", "pay_2"))
        self.assertEqual(self._activate("cs_2", "pay_2")["tier"], "free")

    def test_cancel_on_free_keeps_used_scans(self):
        compliance_db.subscription_snapshot("owner")
        compliance_db._reference("/paperpilot/subscriptions/owner").update({"scans_used": 3})
        self.assertEqual(compliance_db.cancel_user_subscription("owner")["used"], 3)

    def test_free_overshoot_does_not_refill_allowance(self):
        compliance_db.subscription_snapshot("owner")
        compliance_db._reference("/paperpilot/subscriptions/owner").update({"scans_used": 5})
        snap = compliance_db.subscription_snapshot("owner")
        self.assertEqual(snap["tier"], "free")
        self.assertEqual(snap["remaining"], 0)

    def test_running_scans_count_toward_limit(self):
        mechanics = self.mechanics()
        version, _ = compliance_db.create_version(
            "owner", "Paper", mechanics["id"], "paper.docx", "docx", "text", self.parsed(), None,
        )
        compliance_db.subscription_snapshot("owner")
        compliance_db._reference("/paperpilot/subscriptions/owner").update({"scans_used": 2})
        compliance_db._reference("/paperpilot/compliance_scans/owner/s1").set(
            {"id": "s1", "status": "running", "created_at": compliance_db._iso(compliance_db._now())}
        )
        with self.assertRaises(compliance_db.UpgradeRequired):
            compliance_db.check_scan_eligibility("owner", version["manuscript_id"], version["id"])

    def _version_with_scan(self, mechanics, filename):
        version, _ = compliance_db.create_version(
            "owner", "Thesis", mechanics["id"], filename, "docx", "text", self.parsed(), None,
        )
        scan = compliance_db.persist_scan(
            "owner", version["manuscript_id"], version["id"], mechanics["id"], 90.0, [],
            [{"section": "General", "formatting_score": 90.0, "issue_count": 0, "issues": []}],
        )
        return version, scan

    def test_deleting_a_version_keeps_numbers_and_records_its_scans(self):
        self._activate()
        mechanics = self.mechanics()
        v1, s1 = self._version_with_scan(mechanics, "a.docx")
        v2, s2 = self._version_with_scan(mechanics, "b.docx")
        v3, s3 = self._version_with_scan(mechanics, "c.docx")
        mid = v1["manuscript_id"]

        self.assertTrue(compliance_db.delete_version("owner", mid, v2["id"]))
        self.assertIsNone(compliance_db.get_scan("owner", s2["id"]))
        self.assertEqual(compliance_db.deleted_scan_ids("owner"), [s2["id"]])
        summaries = compliance_db.list_scan_summaries("owner")
        self.assertEqual([s["version_number"] for s in summaries], [1, 3])
        self.assertTrue(compliance_db.is_current_version("owner", mid, v3["id"]))

        v4, _ = self._version_with_scan(mechanics, "d.docx")
        self.assertEqual(v4["version_number"], 4)

        self.assertTrue(compliance_db.delete_version("owner", "other-manuscript", v4["id"]))
        self.assertTrue(compliance_db.is_current_version("owner", mid, v3["id"]))
        self.assertFalse(compliance_db.delete_version("owner", mid, v4["id"]))

    def test_deleting_the_last_version_removes_the_manuscript(self):
        mechanics = self.mechanics()
        v1, s1 = self._version_with_scan(mechanics, "a.docx")
        self.assertTrue(compliance_db.delete_version("owner", v1["manuscript_id"], v1["id"]))
        self.assertEqual(compliance_db.list_manuscripts("owner"), [])
        self.assertEqual(compliance_db.deleted_scan_ids("owner"), [s1["id"]])

    def test_version_with_running_scan_cannot_be_deleted(self):
        mechanics = self.mechanics()
        version, _ = compliance_db.create_version(
            "owner", "Paper", mechanics["id"], "paper.docx", "docx", "text", self.parsed(), None,
        )
        compliance_db._reference("/paperpilot/compliance_scans/owner/s9").set(
            {"id": "s9", "status": "running", "manuscript_version_id": version["id"]}
        )
        with self.assertRaises(compliance_db.VersionScanRunning):
            compliance_db.delete_version("owner", version["manuscript_id"], version["id"])


class OtpTests(FirebaseTestCase):
    def test_new_code_invalidates_old_and_challenge_is_single_use(self):
        with patch.object(otp.time, "time", return_value=1000):
            otp.store_otp("USER@example.com", "verify_email", "111111")
        with patch.object(otp.time, "time", return_value=1061):
            otp.store_otp("user@example.com", "verify_email", "222222")
            with self.assertRaisesRegex(ValueError, "Incorrect code"):
                otp.verify_otp("user@example.com", "verify_email", "111111")
            token = otp.verify_otp("user@example.com", "verify_email", "222222")
            otp.consume_challenge(token, "USER@example.com", "verify_email")
            with self.assertRaisesRegex(ValueError, "expired or was already used"):
                otp.consume_challenge(token, "user@example.com", "verify_email")

    def test_attempt_limit_removes_challenge(self):
        with patch.object(otp.time, "time", return_value=2000):
            otp.store_otp("user@example.com", "reset_password", "123456")
            for remaining in (4, 3, 2, 1):
                with self.assertRaisesRegex(ValueError, f"{remaining} attempt"):
                    otp.verify_otp("user@example.com", "reset_password", "000000")
            with self.assertRaisesRegex(ValueError, "Too many incorrect attempts"):
                otp.verify_otp("user@example.com", "reset_password", "000000")
            with self.assertRaisesRegex(ValueError, "No active code"):
                otp.verify_otp("user@example.com", "reset_password", "123456")

    def test_simultaneous_requests_keep_codes_isolated_per_email(self):
        with patch.object(otp.time, "time", return_value=3000):
            otp.store_otp("vince@example.com", "reset_password", "111111")
            otp.store_otp("jemusu@example.com", "reset_password", "222222")
            with self.assertRaisesRegex(ValueError, "Incorrect code"):
                otp.verify_otp("vince@example.com", "reset_password", "222222")
            with self.assertRaisesRegex(ValueError, "Incorrect code"):
                otp.verify_otp("jemusu@example.com", "reset_password", "111111")
            vince_token = otp.verify_otp("vince@example.com", "reset_password", "111111")
            with self.assertRaisesRegex(ValueError, "does not match"):
                otp.consume_challenge(vince_token, "jemusu@example.com", "reset_password")
            jemusu_token = otp.verify_otp("jemusu@example.com", "reset_password", "222222")
            otp.consume_challenge(jemusu_token, "jemusu@example.com", "reset_password")


class OtpEmailTests(unittest.TestCase):
    def test_each_code_is_delivered_only_to_its_requester(self):
        sent = []

        def resend(to_email, subject, text, html):
            if to_email != "owner@example.com":
                raise RuntimeError("You can only send testing emails to your own email address")
            sent.append(("resend", to_email, subject))

        def smtp(to_email, subject, text, html):
            sent.append(("smtp", to_email, subject))

        with (
            patch.object(emailer.settings, "resend_api_key", "re_test"),
            patch.object(emailer.settings, "smtp_host", "smtp.example.com"),
            patch.object(emailer, "_send_via_resend", side_effect=resend),
            patch.object(emailer, "_send_via_smtp", side_effect=smtp),
        ):
            emailer.send_otp_email("owner@example.com", "111111", "reset_password")
            emailer.send_otp_email("jemusu@example.com", "222222", "reset_password")

        self.assertEqual(
            sent,
            [
                ("resend", "owner@example.com", "111111 is your PaperPilot password reset code"),
                ("smtp", "jemusu@example.com", "222222 is your PaperPilot password reset code"),
            ],
        )


if __name__ == "__main__":
    unittest.main()
