#!/usr/bin/env python3
"""
Reproduction harness for Acknowledgement-page heading/measurement bugs.

Usage (from PaperPilot-main/):
  python scripts/repro_scan.py
  python scripts/repro_scan.py --manuscript path/to/file.docx
  python scripts/repro_scan.py --mechanics path/to/rules.json
  python scripts/repro_scan.py --copy api|ml|both
"""

from __future__ import annotations

import argparse
import io
import json
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def _make_ack_docx(path: Path, *, use_heading1_style_on_title: bool = False, all_caps_body: bool = False) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
    from docx.shared import Inches, Pt

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    # Page 1: long title (all caps), authors, then page break
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(
        "PAPERPILOT: AN AI-POWERED ACADEMIC DOCUMENT ANALYSIS AND COMPLIANCE CHECKING SYSTEM"
    )
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.bold = True
    if use_heading1_style_on_title:
        title.style = doc.styles["Heading 1"]
        # Keep explicit 14 pt on the run so Word shows 14 even under Heading 1.
        run.font.size = Pt(14)

    authors = doc.add_paragraph()
    authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ar = authors.add_run("Juan Dela Cruz")
    ar.font.name = "Times New Roman"
    ar.font.size = Pt(11)

    # Force page 2 via hard page break
    pb = doc.add_paragraph()
    pb.add_run().add_break(WD_BREAK.PAGE)

    # Page 2: ACKNOWLEDGEMENT + multi-line indented body at 11 pt (Normal, not Heading)
    ack = doc.add_paragraph()
    ack.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ack_run = ack.add_run("ACKNOWLEDGEMENT")
    ack_run.font.name = "Times New Roman"
    ack_run.font.size = Pt(14)
    ack_run.bold = True

    bodies = [
        (
            "I would like to express my sincere gratitude to the panel members whose insightful "
            "suggestions and helpful comments contributed greatly to the refinement and quality of this study. "
            "Their guidance shaped the direction of the research from proposal to final manuscript."
        ),
        (
            "Special thanks are also extended to the faculty and staff of the College of Engineering "
            "for providing laboratory access, technical advice, and continuous encouragement throughout "
            "the duration of this investigation."
        ),
        (
            "Finally, the author acknowledges family and friends for their patience, understanding, "
            "and unwavering support during the long hours spent writing, revising, and preparing this work."
        ),
        (
            "Without these people the completion of this research study would not have been possible, "
            "and any remaining errors are the sole responsibility of the author alone."
        ),
        (
            "May this document serve as a reminder that scholarly work is always a shared effort, "
            "built on the generosity of mentors, peers, and communities who care about careful inquiry."
        ),
    ]
    if all_caps_body:
        bodies = [b.upper() for b in bodies]

    for text in bodies:
        para = doc.add_paragraph()
        para.paragraph_format.first_line_indent = Inches(0.5)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = para.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(11)
        r.bold = False

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def _load_copy(which: str):
    """Import parse_document / run_compliance_scan from api or machinelearning."""
    if which == "api":
        sys.path.insert(0, str(ROOT / "api"))
    elif which == "ml":
        sys.path.insert(0, str(ROOT / "machinelearning"))
    else:
        raise ValueError(which)
    # Fresh imports for each copy
    for name in list(sys.modules):
        if name == "app" or name.startswith("app."):
            del sys.modules[name]
    from app.documents import parse_document, normalize_mechanics_rules  # noqa: WPS433
    from app.compliance import (  # noqa: WPS433
        run_compliance_scan,
        _document_units,
        _annotate_roles,
        _expected_size_for_role,
        _heading_sizes,
        _defined_heading_levels,
        _looks_like_heading_line,
        _is_structural_heading_text,
        _dominant_size,
        HEADING_STYLE_RE,
    )

    return {
        "parse_document": parse_document,
        "normalize_mechanics_rules": normalize_mechanics_rules,
        "run_compliance_scan": run_compliance_scan,
        "_document_units": _document_units,
        "_annotate_roles": _annotate_roles,
        "_expected_size_for_role": _expected_size_for_role,
        "_heading_sizes": _heading_sizes,
        "_defined_heading_levels": _defined_heading_levels,
        "_looks_like_heading_line": _looks_like_heading_line,
        "_is_structural_heading_text": _is_structural_heading_text,
        "_dominant_size": _dominant_size,
        "HEADING_STYLE_RE": HEADING_STYLE_RE,
    }


def _role_why(unit: dict, font: dict, mods) -> str:
    text = (unit.get("text") or "").strip()
    style = str(unit.get("style") or "").lower()
    defined = mods["_defined_heading_levels"](font)
    heading_sizes = mods["_heading_sizes"](font, include_legacy=False)
    if unit.get("role") == "body":
        reasons = []
        if unit.get("soft_wrap_continuation"):
            reasons.append("soft_wrap_continuation=>not heading_line")
        if not mods["_looks_like_heading_line"](unit, text):
            reasons.append("failed _looks_like_heading_line")
        elif not mods["_is_structural_heading_text"](text) and not mods["HEADING_STYLE_RE"].search(style):
            reasons.append("not structural / not Heading style")
        return "body (" + (", ".join(reasons) or "default") + ")"
    parts = [f"role={unit.get('role')}"]
    if unit.get("soft_wrap_continuation"):
        parts.append("soft_wrap_continuation=True")
    parts.append(f"looks_like_heading={mods['_looks_like_heading_line'](unit, text)}")
    parts.append(f"structural={mods['_is_structural_heading_text'](text)}")
    style_m = mods["HEADING_STYLE_RE"].search(style)
    parts.append(f"style={style!r}" + (f"=>heading{style_m.group(1)}" if style_m else ""))
    parts.append(f"defined={sorted(defined)}")
    parts.append(f"heading_sizes={heading_sizes}")
    return "; ".join(parts)


def _heading_key_paths(rules: dict) -> list[str]:
    found = []

    def walk(obj, prefix=""):
        if isinstance(obj, dict):
            for k, v in obj.items():
                path = f"{prefix}.{k}" if prefix else k
                if "heading" in str(k).lower() or (isinstance(k, str) and k.endswith("_size")):
                    found.append(f"{path} = {v!r}")
                walk(v, path)
        elif isinstance(obj, list):
            for i, item in enumerate(obj):
                walk(item, f"{prefix}[{i}]")

    walk(rules)
    return found


def run_one(which: str, manuscript: Path, rules: dict, label: str, pages=(0, 1)) -> dict:
    mods = _load_copy(which)
    data = manuscript.read_bytes()
    parsed = mods["parse_document"](data, manuscript.suffix.lstrip(".").lower())
    normalized = mods["normalize_mechanics_rules"](rules)
    font = normalized.get("font") or {}

    units = mods["_document_units"](parsed, normalized)
    mods["_annotate_roles"](units, normalized)

    print("\n" + "=" * 88)
    print(f"COPY={which}  LABEL={label}")
    print(f"manuscript={manuscript}")
    print(f"normalized font keys: {sorted(font.keys())}")
    print(f"heading key paths: {_heading_key_paths(normalized) or '(none)'}")
    print(f"_heading_sizes={mods['_heading_sizes'](font)}")
    print(f"_defined_heading_levels={sorted(mods['_defined_heading_levels'](font))}")

    # Raw paragraph run sizes (parser output)
    print("\n--- Parsed paragraph run sizes (first 12 body paras) ---")
    shown = 0
    for para in parsed.get("paragraphs") or []:
        if (para.get("source") or "body") != "body":
            continue
        sizes = [(r.get("size"), (r.get("text") or "")[:40]) for r in (para.get("runs") or [])]
        print(
            f"  p{para.get('paragraph_index')} style={para.get('style')!r} "
            f"text={(para.get('text') or '')[:50]!r} runs={sizes}"
        )
        shown += 1
        if shown >= 12:
            break

    print(f"\n--- Units on pages {list(pages)} ---")
    for unit in units:
        if unit.get("page_index") not in pages:
            continue
        text = (unit.get("text") or "").strip()
        if not text:
            continue
        role = unit.get("role")
        measured = unit.get("dominant_size")
        if measured is None:
            measured = mods["_dominant_size"](unit)
        line_run_sizes = [r.get("size") for r in (unit.get("line_runs") or [])]
        para_run_sizes = [r.get("size") for r in (unit.get("runs") or []) if str(r.get("text") or "").strip()]
        expected = mods["_expected_size_for_role"](role, font)
        expected_src = (
            f"font.{role}_size"
            if role and str(role).startswith("heading") and expected is not None
            else (
                "font.body_size|heading3_content_size"
                if role in {"body", "caption"} and expected is not None
                else "(unset)"
            )
        )
        print(
            f"  p{unit['page_index'] + 1} L{unit['line_index']}: role={role} "
            f"measured={measured} line_run_sizes={line_run_sizes} "
            f"para_run_sizes={para_run_sizes[:4]} expected={expected} src={expected_src}"
        )
        print(f"    excerpt={text[:90]!r}")
        print(f"    why: {_role_why(unit, font, mods)}")
        print(
            f"    soft_wrap_cont={unit.get('soft_wrap_continuation')} "
            f"style={unit.get('style')!r} indent={(unit.get('formatting') or {}).get('first_line_indent_inches')}"
        )

    result = mods["run_compliance_scan"](parsed, normalized, "free")
    issues = result.get("issues") or []
    types = sorted({i.get("issue_type") for i in issues})
    heading_issues = [
        i
        for i in issues
        if "heading" in str(i.get("issue_type") or "").lower()
        or "Heading" in str(i.get("title") or "")
    ]
    print(f"\n--- Scan summary ---")
    print(f"issue_count={len(issues)} types={types}")
    print(f"heading_related_count={len(heading_issues)}")
    for i in heading_issues[:8]:
        print(
            f"  ISSUE {i.get('issue_type')}: {i.get('title')}\n"
            f"    explanation={i.get('explanation')}\n"
            f"    locs={len(i.get('locations') or [])}"
        )
        for loc in (i.get("locations") or [])[:3]:
            print(
                f"      loc p{loc.get('page')} L{loc.get('line')}: {(loc.get('excerpt') or '')[:80]!r}"
            )
    scoring = result.get("scoring") or {}
    print(
        f"scoring keys sample: overall={scoring.get('overall_score')} "
        f"categories={[(c.get('category'), c.get('units_checked'), c.get('units_failed')) for c in scoring.get('categories') or [] if c.get('category')=='Fonts']}"
    )
    return {
        "which": which,
        "label": label,
        "issue_count": len(issues),
        "types": types,
        "heading_related_count": len(heading_issues),
        "issues": issues,
        "normalized": normalized,
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--manuscript", type=Path, default=None)
    parser.add_argument("--mechanics", type=Path, default=None)
    parser.add_argument("--copy", choices=("api", "ml", "both"), default="both")
    parser.add_argument("--out-dir", type=Path, default=ROOT / "testdata" / "repro_ack")
    args = parser.parse_args()

    out = args.out_dir
    out.mkdir(parents=True, exist_ok=True)

    if args.manuscript:
        manuscript = args.manuscript
    else:
        manuscript = out / "ack_page_fixture.docx"
        _make_ack_docx(manuscript, use_heading1_style_on_title=False, all_caps_body=False)
        # Also write variants for later cases
        _make_ack_docx(out / "ack_heading1_style_title.docx", use_heading1_style_on_title=True)
        _make_ack_docx(out / "ack_all_caps_body.docx", all_caps_body=True)

    if args.mechanics:
        rules = json.loads(args.mechanics.read_text(encoding="utf-8"))
        cases = [("user_mechanics", rules)]
    else:
        cases = [
            (
                "NO_heading_rules",
                {
                    "font": {
                        "families": ["Times New Roman"],
                        "body_size": 11,
                        "sizes_points": [11.0],
                        "color": "Black/Automatic",
                    }
                },
            ),
            (
                "H1_explicit_16",
                {
                    "font": {
                        "families": ["Times New Roman"],
                        "body_size": 11,
                        "heading1_size": 16,
                        "heading_levels": 1,
                    }
                },
            ),
            (
                "phantom_H1_40",
                {
                    "font": {
                        "families": ["Times New Roman"],
                        "body_size": 11,
                        "heading1_size": 40.0,
                        "sizes_points": [11.0],
                    }
                },
            ),
            (
                "list_H1_11",
                {
                    "font": {
                        "families": ["Times New Roman"],
                        "body_size": 11,
                        "heading1_size": [11.0],
                        "sizes_points": [11.0],
                    }
                },
            ),
        ]
        # Persist for inspection
        for name, rules in cases:
            (out / f"mechanics_{name}.json").write_text(json.dumps(rules, indent=2), encoding="utf-8")

    copies = ["api", "ml"] if args.copy == "both" else [args.copy]
    summaries = []
    for which in copies:
        for label, rules in cases:
            summaries.append(run_one(which, manuscript, rules, label))

    print("\n" + "=" * 88)
    print("SUMMARY TABLE")
    print(f"{'copy':<4} {'label':<22} {'issues':>6} {'heading':>8} types")
    for s in summaries:
        print(
            f"{s['which']:<4} {s['label']:<22} {s['issue_count']:>6} "
            f"{s['heading_related_count']:>8} {s['types']}"
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
